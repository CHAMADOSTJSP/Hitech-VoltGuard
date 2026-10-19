import os
import io
import json
import random
import hashlib
import secrets
import hmac
import socket
import ipaddress
from urllib.parse import urlparse
import httpx
from datetime import date, datetime, timedelta, timezone
from typing import List, Optional
from fastapi import FastAPI, Depends, HTTPException, status, WebSocket, WebSocketDisconnect, Request, Header
from fastapi.security import OAuth2PasswordBearer
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from pydantic import BaseModel
import jwt
from sqlalchemy import create_engine, Column, Integer, String, Float, DateTime, Boolean, Index, Text, text, or_
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker, Session
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from slowapi.util import get_remote_address

# Configurações de Segurança e JWT
# Configure SECRET_KEY in the environment for deployments. The fallback is intentionally
# ephemeral: tokens are invalidated when the process restarts if no key is configured.
SECRET_KEY = os.getenv("SECRET_KEY")
if not SECRET_KEY and os.getenv("RENDER"):
    raise RuntimeError("Configure SECRET_KEY nos Environment Variables do Render.")
SECRET_KEY = SECRET_KEY or secrets.token_urlsafe(48)
if len(SECRET_KEY) < 32:
    raise RuntimeError("SECRET_KEY deve ter pelo menos 32 caracteres.")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="api/token")

# Configuração da Base de Dados SQLAlchemy
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./voltguard.db")
# Render pode fornecer URLs PostgreSQL com o esquema postgres://. SQLAlchemy 2
# usa o driver psycopg explicitamente para evitar depender de psycopg2.
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql+psycopg://", 1)
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)
engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {},
    pool_pre_ping=True,
)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

# --- MODELOS DA BASE DE DADOS ---
class PanelModel(Base):
    __tablename__ = "panels"
    id = Column(String, primary_key=True, index=True)
    tenant_id = Column(String, default="HITECH-DEFAULT", index=True)
    name = Column(String, nullable=False)
    location = Column(String, nullable=False)
    voltage = Column(String, nullable=False)
    current = Column(Float, default=0.0)
    temperature = Column(Float, default=0.0)
    temperature_measured = Column(Boolean, nullable=False, default=True)
    voltage_v = Column(Float, nullable=True)
    active_power_kw = Column(Float, nullable=True)
    active_energy_kwh = Column(Float, nullable=True)
    frequency_hz = Column(Float, nullable=True)
    power_factor = Column(Float, nullable=True)
    gateway_qr = Column(String, nullable=True, index=True)
    status = Column(String, default="NORMAL")
    component = Column(String, default="Disjuntor Principal")
    predictive_failure_hours = Column(Integer, default=120)
    recommendation = Column(String, default="Operação nominal estável.")
    esg_waste_kwh = Column(Float, default=1.2)
    assigned_to = Column(String, nullable=True)

class PanelHistoryModel(Base):
    __tablename__ = "panel_history"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    panel_id = Column(String, index=True, nullable=False)
    tenant_id = Column(String, index=True, nullable=False, default="HITECH-DEFAULT")
    timestamp = Column(String, nullable=False)
    temperature = Column(Float, nullable=False)
    # Nullable for backward compatibility with historical rows created before v0.5.1.
    current = Column(Float, nullable=True)
    voltage = Column(String, nullable=True)
    temperature_measured = Column(Boolean, nullable=False, default=True)
    voltage_v = Column(Float, nullable=True)
    active_power_kw = Column(Float, nullable=True)
    active_energy_kwh = Column(Float, nullable=True)
    frequency_hz = Column(Float, nullable=True)
    power_factor = Column(Float, nullable=True)

class AuditLogModel(Base):
    __tablename__ = "audit_logs"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, default="HITECH-DEFAULT", index=True)
    timestamp = Column(DateTime, default=datetime.utcnow)
    user_name = Column(String, nullable=False)
    role = Column(String, nullable=False)
    action = Column(String, nullable=False)
    details = Column(String, nullable=False)

class WorkOrderHistoryModel(Base):
    __tablename__ = "work_order_history"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, nullable=False, index=True)
    work_order_id = Column(Integer, nullable=False, index=True)
    timestamp = Column(DateTime, default=datetime.utcnow, nullable=False, index=True)
    user_name = Column(String, nullable=False)
    role = Column(String, nullable=False)
    action = Column(String, nullable=False)
    details = Column(Text, nullable=False)
    resolution = Column(Text, nullable=True)

class WorkOrderChecklistItemModel(Base):
    __tablename__ = "work_order_checklist_items"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, nullable=False, index=True)
    work_order_id = Column(Integer, nullable=False, index=True)
    label = Column(String(240), nullable=False)
    is_completed = Column(Boolean, default=False, nullable=False)
    completed_at = Column(DateTime, nullable=True)
    completed_by = Column(String, nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)

class WorkOrderChecklistTemplateModel(Base):
    __tablename__ = "work_order_checklist_templates"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, nullable=False, index=True)
    name = Column(String(120), nullable=False)
    items_json = Column(Text, nullable=False)
    created_by = Column(String, nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    __table_args__ = (Index("uq_work_order_checklist_template_tenant_name", tenant_id, name, unique=True),)

class TenantSubscriptionModel(Base):
    __tablename__ = "subscriptions"
    tenant_id = Column(String, primary_key=True, index=True)
    plan_name = Column(String, default="Enterprise")
    status = Column(String, default="ACTIVE")
    expires_at = Column(DateTime, default=datetime.utcnow() + timedelta(days=365))
    max_panels = Column(Integer, default=25)
    webhook_url = Column(String, nullable=True, default="https://httpbin.org/post")

class TenantRuleModel(Base):
    __tablename__ = "tenant_rules"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    max_temperature = Column(Float, default=85.0)
    max_current = Column(Float, default=500.0)

class PanelAlarmRuleModel(Base):
    __tablename__ = "panel_alarm_rules"
    id = Column(Integer, primary_key=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    panel_id = Column(String, index=True, nullable=False)
    max_temperature = Column(Float, nullable=False, default=85.0)
    max_current = Column(Float, nullable=False, default=500.0)
    temperature_hysteresis = Column(Float, nullable=False, default=3.0)
    current_hysteresis = Column(Float, nullable=False, default=10.0)
    updated_at = Column(DateTime, default=datetime.utcnow, nullable=False)


class WorkOrderModel(Base):
    __tablename__ = "work_orders"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, default="HITECH-DEFAULT", index=True)
    alarm_id = Column(Integer, nullable=True, index=True)
    panel_id = Column(String, nullable=False)
    title = Column(String, nullable=False)
    priority = Column(String, default="ALTA")
    status = Column(String, default="ABERTA")
    assigned_technician = Column(String, nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow)
    description = Column(String, nullable=True)
    due_at = Column(DateTime, nullable=True)
    started_at = Column(DateTime, nullable=True)
    completed_at = Column(DateTime, nullable=True)
    resolution = Column(String, nullable=True)


Index("uq_work_orders_tenant_alarm", WorkOrderModel.tenant_id, WorkOrderModel.alarm_id, unique=True)

class InventoryModel(Base):
    __tablename__ = "inventory"
    id = Column(Integer, primary_key=True, index=True, autoincrement=True)
    tenant_id = Column(String, default="HITECH-DEFAULT", index=True)
    part_name = Column(String, nullable=False)
    category = Column(String, default="Elétrica / Manutenção")
    quantity_in_stock = Column(Integer, default=10)
    min_threshold = Column(Integer, default=3)
    unit_cost = Column(Float, default=150.0)

class AlarmModel(Base):
    __tablename__ = "alarms"
    id = Column(Integer, primary_key=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    panel_id = Column(String, index=True, nullable=False)
    alarm_type = Column(String, nullable=False)
    severity = Column(String, default="WARNING", nullable=False)
    message = Column(String, nullable=False)
    status = Column(String, default="ACTIVE", nullable=False)  # ACTIVE, ACKNOWLEDGED, CLEARED
    first_triggered_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    last_triggered_at = Column(DateTime, default=datetime.utcnow, nullable=False)
    acknowledged_at = Column(DateTime, nullable=True)
    acknowledged_by = Column(String, nullable=True)
    cleared_at = Column(DateTime, nullable=True)


class AssetModel(Base):
    __tablename__ = "assets"
    id = Column(Integer, primary_key=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    asset_tag = Column(String, nullable=False, index=True)
    name = Column(String, nullable=False)
    asset_type = Column(String, default="PAINEL_ELETRICO")
    location = Column(String, default="")
    manufacturer = Column(String, default="")
    model = Column(String, default="")
    serial_number = Column(String, default="")
    panel_id = Column(String, nullable=True, index=True)
    status = Column(String, default="ATIVO")
    notes = Column(String, default="")
    created_at = Column(DateTime, default=datetime.utcnow)

class MaintenancePlanModel(Base):
    __tablename__ = "maintenance_plans"
    id = Column(Integer, primary_key=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    asset_id = Column(Integer, index=True, nullable=False)
    title = Column(String, nullable=False)
    checklist = Column(String, default="Inspeção visual;Verificar conexões;Registrar observações")
    interval_days = Column(Integer, default=30)
    next_due_at = Column(DateTime, nullable=False)
    active = Column(Integer, default=1)
    created_at = Column(DateTime, default=datetime.utcnow)

class MaintenanceTaskModel(Base):
    __tablename__ = "maintenance_tasks"
    id = Column(Integer, primary_key=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    plan_id = Column(Integer, index=True, nullable=True)
    asset_id = Column(Integer, index=True, nullable=False)
    title = Column(String, nullable=False)
    due_at = Column(DateTime, nullable=False)
    status = Column(String, default="PENDENTE")
    assigned_to = Column(String, nullable=True)
    notes = Column(String, default="")
    created_at = Column(DateTime, default=datetime.utcnow)
    completed_at = Column(DateTime, nullable=True)

class DeviceCredentialModel(Base):
    __tablename__ = "device_credentials"
    id = Column(Integer, primary_key=True, autoincrement=True)
    tenant_id = Column(String, index=True, nullable=False)
    panel_id = Column(String, index=True, nullable=False)
    name = Column(String, nullable=False)
    token_hash = Column(String, unique=True, index=True, nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow)
    last_used_at = Column(DateTime, nullable=True)
    revoked_at = Column(DateTime, nullable=True)

class UserModel(Base):
    __tablename__ = "users"
    username = Column(String, primary_key=True, index=True)
    password = Column(String, nullable=False)
    role = Column(String, default="CEO")
    tenant_id = Column(String, default="HITECH-DEFAULT")
    plan = Column(String, default="Enterprise")

def validate_webhook_url(value: str) -> str:
    """Permite apenas HTTPS público, sem credenciais, IPs privados ou hosts locais."""
    if not value:
        return ""
    if len(value) > 2048:
        raise HTTPException(status_code=422, detail="URL de webhook demasiado longa.")
    try:
        parsed = urlparse(value)
        port = parsed.port or 443
    except ValueError:
        raise HTTPException(status_code=422, detail="URL de webhook inválida.")
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise HTTPException(status_code=422, detail="O webhook deve ser uma URL HTTPS pública válida.")
    host = parsed.hostname.rstrip(".").lower()
    if host in {"localhost", "localhost.localdomain"} or host.endswith((".local", ".internal", ".localhost")):
        raise HTTPException(status_code=422, detail="Host local não permitido para webhook.")
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)}
    except OSError:
        raise HTTPException(status_code=422, detail="Não foi possível resolver o host do webhook.")
    if not addresses:
        raise HTTPException(status_code=422, detail="Host do webhook sem endereço válido.")
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if not ip.is_global:
            raise HTTPException(status_code=422, detail="Webhooks para endereços privados ou especiais não são permitidos.")
    return value

# Instanciação primária da aplicação FastAPI
app = FastAPI(
    title="HITECH VoltGuard SCADA Enterprise",
    description="MVP SaaS de demonstração para monitoramento e manutenção. Não certificado para proteção ou comando industrial.",
    version="0.39.0",
)
limiter = Limiter(key_func=get_remote_address, default_limits=[])
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

@app.middleware("http")
async def security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(self), microphone=(), geolocation=()"
    if request.url.path.startswith("/api/") or request.url.path == "/health":
        response.headers["Cache-Control"] = "no-store"
    response.headers["Cross-Origin-Resource-Policy"] = "same-origin"
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response

@app.get("/health", tags=["Operations"])
def health_check():
    try:
        with engine.connect() as connection:
            connection.execute(text("SELECT 1"))
    except Exception:
        raise HTTPException(status_code=503, detail="Banco de dados indisponível.")
    return {"status": "ok", "service": "voltguard", "mode": "simulation" if os.getenv("SIMULATE_TELEMETRY", "0").lower() in {"1", "true", "yes"} else "stored-data"}


# Gestor de Conexões WebSocket
class ConnectionManager:
    def __init__(self):
        self.active_connections: dict[WebSocket, str] = {}

    async def connect(self, websocket: WebSocket, tenant_id: str):
        await websocket.accept()
        self.active_connections[websocket] = tenant_id

    def disconnect(self, websocket: WebSocket):
        self.active_connections.pop(websocket, None)

    async def broadcast(self, message: dict, tenant_id: str):
        stale = []
        for connection, connection_tenant in list(self.active_connections.items()):
            if connection_tenant != tenant_id:
                continue
            try:
                await connection.send_json(message)
            except Exception:
                stale.append(connection)
        for connection in stale:
            self.disconnect(connection)

manager = ConnectionManager()

# Funções Utilitárias de Criptografia
def hash_password(password: str) -> str:
    """Hash adaptativo com scrypt, usando apenas a biblioteca padrão."""
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1)
    return "scrypt$%s$%s" % (salt.hex(), digest.hex())

def verify_password(plain_password: str, stored_password: str) -> bool:
    if not stored_password:
        return False
    try:
        if stored_password.startswith("scrypt$"):
            _, salt_hex, digest_hex = stored_password.split("$", 2)
            actual = hashlib.scrypt(plain_password.encode("utf-8"), salt=bytes.fromhex(salt_hex), n=2**14, r=8, p=1)
            return hmac.compare_digest(actual.hex(), digest_hex)
        # Compatibilidade temporária com hashes SHA-256 antigos; o login migra o hash.
        legacy = hashlib.sha256(plain_password.encode("utf-8")).hexdigest()
        return hmac.compare_digest(legacy, stored_password)
    except (ValueError, TypeError):
        return False

# Gestão de Sessões e Inicialização da Base de Dados
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def utc_iso(value: Optional[datetime]) -> Optional[str]:
    if value is None:
        return None
    if value.tzinfo is not None:
        value = value.astimezone(timezone.utc).replace(tzinfo=None)
    return value.isoformat(timespec="seconds") + "Z"

def utc_naive(value: Optional[datetime]) -> Optional[datetime]:
    """Normalize API datetimes to naive UTC for the existing DateTime columns."""
    if value is None:
        return None
    if value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value

def init_db():
    db = SessionLocal()
    try:
        # Compatibility migration for deployments created before panel_history.tenant_id existed.
        # This is deliberately narrow; structural changes beyond this must use a formal migration.
        with engine.begin() as connection:
            inspector = __import__("sqlalchemy").inspect(connection)
            if "panel_history" in inspector.get_table_names():
                columns = {col["name"] for col in inspector.get_columns("panel_history")}
                if "tenant_id" not in columns:
                    connection.execute(text("ALTER TABLE panel_history ADD COLUMN tenant_id VARCHAR"))
                    connection.execute(text("UPDATE panel_history SET tenant_id = COALESCE((SELECT panels.tenant_id FROM panels WHERE panels.id = panel_history.panel_id), 'ORPHANED') WHERE tenant_id IS NULL"))
                # Backward-compatible additions: old history rows retain NULL for unavailable measurements.
                if "current" not in columns:
                    connection.execute(text("ALTER TABLE panel_history ADD COLUMN current FLOAT"))
                if "voltage" not in columns:
                    connection.execute(text("ALTER TABLE panel_history ADD COLUMN voltage VARCHAR"))
                history_columns = {
                    "temperature_measured": "BOOLEAN NOT NULL DEFAULT 1",
                    "voltage_v": "FLOAT",
                    "active_power_kw": "FLOAT",
                    "active_energy_kwh": "FLOAT",
                    "frequency_hz": "FLOAT",
                    "power_factor": "FLOAT",
                }
                for column_name, column_type in history_columns.items():
                    if column_name not in columns:
                        connection.execute(text(f"ALTER TABLE panel_history ADD COLUMN {column_name} {column_type}"))
            if "panels" in inspector.get_table_names():
                panel_columns = {col["name"] for col in inspector.get_columns("panels")}
                if "temperature_measured" not in panel_columns:
                    connection.execute(text("ALTER TABLE panels ADD COLUMN temperature_measured BOOLEAN NOT NULL DEFAULT 1"))
                telemetry_columns = {
                    "voltage_v": "FLOAT",
                    "active_power_kw": "FLOAT",
                    "active_energy_kwh": "FLOAT",
                    "frequency_hz": "FLOAT",
                    "power_factor": "FLOAT",
                }
                for column_name, column_type in telemetry_columns.items():
                    if column_name not in panel_columns:
                        connection.execute(text(f"ALTER TABLE panels ADD COLUMN {column_name} {column_type}"))
            if "work_orders" in inspector.get_table_names():
                columns = {col["name"] for col in inspector.get_columns("work_orders")}
                compatibility_columns = {
                    "alarm_id": "INTEGER",
                    "due_at": "TIMESTAMP",
                    "started_at": "TIMESTAMP",
                    "completed_at": "TIMESTAMP",
                    "resolution": "TEXT",
                }
                for column_name, column_type in compatibility_columns.items():
                    if column_name not in columns:
                        connection.execute(text(f"ALTER TABLE work_orders ADD COLUMN {column_name} {column_type}"))
                connection.execute(text(
                    "CREATE UNIQUE INDEX IF NOT EXISTS uq_work_orders_tenant_alarm "
                    "ON work_orders (tenant_id, alarm_id)"
                ))
        Base.metadata.create_all(bind=engine)
        
        # Não criar contas de demonstração com senhas conhecidas. Todas as contas
        # devem ser criadas explicitamente pelo fluxo de cadastro da organização.

        # Fresh installations start empty. The registration flow creates each tenant's
        # subscription, rules, and clearly marked demonstration panels.
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

init_db()

# Esquemas Pydantic
class User(BaseModel):
    username: str
    role: str
    tenant_id: str
    plan: str

class TelemetryData(BaseModel):
    panel_id: str
    temperature: float
    current: float
    voltage: Optional[str] = "380V AC"
    gateway_qr: Optional[str] = None


# Dependência de Autenticação JWT
def get_current_user(token: str = Depends(oauth2_scheme)) -> User:
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username = payload.get("sub")
        role = payload.get("role")
        tenant_id = payload.get("tenant_id")
        plan = payload.get("plan")
        if not username or not role or not tenant_id or not plan:
            raise HTTPException(status_code=401, detail="Token inválido.")
        return User(username=username, role=role, tenant_id=tenant_id, plan=plan)
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Token inválido ou expirado.")

# --- ROTAS DA API ---

@app.get("/logo.jpg")
def get_logo():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    logo_path = os.path.join(base_dir, "logo.jpg")
    if os.path.exists(logo_path):
        return FileResponse(logo_path)
    raise HTTPException(status_code=404, detail="Logo não encontrada.")

@app.post("/api/token")
@limiter.limit("10/minute")
def login(request: Request, credentials: dict, db: Session = Depends(get_db)):
    raw_username = credentials.get("username", "")
    username = raw_username.strip() if isinstance(raw_username, str) else ""
    password = credentials.get("password", "")
    if not isinstance(password, str) or not username or len(username) > 120 or len(password) > 256:
        raise HTTPException(status_code=401, detail="Utilizador ou senha incorretos.")
    
    db_user = db.query(UserModel).filter(UserModel.username == username).first()
    if not db_user or not verify_password(password, db_user.password):
        raise HTTPException(status_code=401, detail="Utilizador ou senha incorretos.")
    if not db_user.password.startswith("scrypt$"):
        db_user.password = hash_password(password)
        db.commit()
    
    token = jwt.encode({"sub": db_user.username, "role": db_user.role, "tenant_id": db_user.tenant_id, "plan": db_user.plan, "exp": datetime.utcnow() + timedelta(days=1)}, SECRET_KEY, algorithm="HS256")
    return {"access_token": token, "token_type": "bearer", "role": db_user.role, "plan": db_user.plan}

@app.post("/api/tenant/register")
@limiter.limit("5/hour")
def register_tenant(request: Request, data: dict, db: Session = Depends(get_db)):
    raw_tenant = data.get("tenant_id", "")
    raw_user = data.get("username", "")
    tenant_id = raw_tenant.upper().strip() if isinstance(raw_tenant, str) else ""
    admin_user = raw_user.strip() if isinstance(raw_user, str) else ""
    password = data.get("password", "")
    
    if not tenant_id or not admin_user or not isinstance(password, str) or not password.strip():
        raise HTTPException(status_code=400, detail="Dados de registo incompletos.")
    if len(password) < 12:
        raise HTTPException(status_code=400, detail="A senha deve ter pelo menos 12 caracteres.")
    if len(password) > 256:
        raise HTTPException(status_code=422, detail="A senha excede o tamanho permitido.")
    if len(tenant_id) > 80 or len(admin_user) > 120:
        raise HTTPException(status_code=400, detail="Identificador ou utilizador demasiado longo.")
    import re
    if not re.fullmatch(r"[A-Z0-9_-]{2,80}", tenant_id) or not re.fullmatch(r"[A-Za-z0-9_.-]{3,120}", admin_user):
        raise HTTPException(status_code=422, detail="Identificador da organização ou utilizador inválido.")
    if db.query(TenantSubscriptionModel).filter(TenantSubscriptionModel.tenant_id == tenant_id).first():
        raise HTTPException(status_code=409, detail="Esta organização já está registada.")
    
    existing = db.query(UserModel).filter(UserModel.username == admin_user).first()
    if existing:
        raise HTTPException(status_code=400, detail="Nome de utilizador já existe.")
        
    new_user = UserModel(username=admin_user, password=hash_password(password), role="CEO", tenant_id=tenant_id, plan="Piloto")
    db.add(new_user)
    db.add(TenantSubscriptionModel(tenant_id=tenant_id, plan_name="Piloto", status="ACTIVE", max_panels=25, webhook_url=None))
    db.add(TenantRuleModel(tenant_id=tenant_id, max_temperature=85.0, max_current=500.0))

    # Dados explicitamente demonstrativos para que uma organização recém-criada
    # possa avaliar a interface sem conectar equipamentos físicos.
    demo_suffix = secrets.token_hex(3).upper()
    demo_panels = [
        ("01", "Alimentação trifásica — DEMO", "Ambiente simulado", "380 V AC", 118.4, 42.5, "NORMAL", "Valores fictícios para demonstração."),
        ("02", "Carga industrial — DEMO", "Ambiente simulado", "380 V AC", 286.7, 68.2, "NORMAL", "Valores fictícios para demonstração."),
        ("03", "Cenário de alerta — DEMO", "Ambiente simulado", "380 V AC", 515.0, 91.0, "CRITICO", "ALERTA SIMULADO: limites demonstrativos excedidos; não representa condição elétrica real."),
    ]
    for suffix, name, location, voltage, current, temperature, panel_status, recommendation in demo_panels:
        panel_id = f"{tenant_id[:55]}-{demo_suffix}-PANEL-{suffix}"
        panel = PanelModel(
            id=panel_id, tenant_id=tenant_id, name=name, location=location,
            voltage=voltage, current=current, temperature=temperature,
            gateway_qr=f"DEMO-{demo_suffix}-{suffix}", status=panel_status,
            component="Simulador VoltGuard", predictive_failure_hours=0,
            recommendation=recommendation, esg_waste_kwh=0.0,
        )
        db.add(panel)
        db.add(PanelHistoryModel(panel_id=panel_id, tenant_id=tenant_id, timestamp=datetime.utcnow().isoformat(timespec="seconds") + "Z", temperature=temperature, current=panel.current, voltage=panel.voltage))

    db.add(AuditLogModel(tenant_id=tenant_id, user_name=admin_user, role="CEO", action="REGISTO_TENANT", details=f"Organização {tenant_id} registada com três painéis demonstrativos."))
    db.commit()
    
    return {"status": "success", "message": f"Organização {tenant_id} criada com sucesso!"}

@app.post("/api/tenant/webhook")
def save_webhook(data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role != "CEO":
        raise HTTPException(status_code=403, detail="Acesso restrito a CEO.")
    webhook_url = validate_webhook_url(str(data.get("webhook_url", "")).strip())
    sub = db.query(TenantSubscriptionModel).filter(TenantSubscriptionModel.tenant_id == current_user.tenant_id).first()
    if not sub:
        raise HTTPException(status_code=404, detail="Assinatura da organização não encontrada.")
    sub.webhook_url = webhook_url or None
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CONFIGURAR_WEBHOOK", details="Webhook atualizado."))
    db.commit()
    return {"status": "success", "message": "Webhook atualizado com sucesso."}

@app.post("/api/tenant/webhook/test")
async def test_webhook(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Envia um evento de teste ao endpoint salvo, sem criar alarmes."""
    if current_user.role != "CEO":
        raise HTTPException(status_code=403, detail="Acesso restrito a CEO.")
    sub = db.query(TenantSubscriptionModel).filter(
        TenantSubscriptionModel.tenant_id == current_user.tenant_id
    ).first()
    if not sub or not sub.webhook_url:
        raise HTTPException(status_code=409, detail="Configure e salve um endpoint antes de enviar o teste.")

    safe_url = validate_webhook_url(sub.webhook_url)
    event_id = secrets.token_urlsafe(12)
    http_status = None
    delivered = False
    try:
        async with httpx.AsyncClient(follow_redirects=False, timeout=5.0) as client:
            response = await client.post(safe_url, json={
                "event": "WEBHOOK_TEST",
                "event_id": event_id,
                "organization_id": current_user.tenant_id,
                "sent_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
            })
        http_status = response.status_code
        delivered = 200 <= http_status < 300
    except Exception:
        # Não devolver exceções de rede, que podem expor dados do endpoint configurado.
        delivered = False

    result = "DELIVERED" if delivered else "FAILED"
    details = f"Teste de webhook {result}" + (f" (HTTP {http_status})." if http_status else ". Sem resposta HTTP confirmada.")
    db.add(AuditLogModel(
        tenant_id=current_user.tenant_id,
        user_name=current_user.username,
        role=current_user.role,
        action="TESTAR_WEBHOOK",
        details=details,
    ))
    db.commit()
    return {
        "status": "success",
        "delivery": result,
        "http_status": http_status,
        "message": "Evento de teste entregue." if delivered else "O endpoint não confirmou a entrega.",
    }

@app.post("/api/tenant/rules")
def update_tenant_rules(data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in ["CEO", "Gerente"]:
        raise HTTPException(status_code=403, detail="Acesso restrito para alterar regras de alarme.")
        
    try:
        max_temp = float(data.get("max_temperature", 85.0))
        max_curr = float(data.get("max_current", 500.0))
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Os limites devem ser numéricos.")
    import math
    if not math.isfinite(max_temp) or not 0 < max_temp <= 250 or not math.isfinite(max_curr) or not 0 < max_curr <= 100000:
        raise HTTPException(status_code=422, detail="Limites fora dos intervalos permitidos.")
    
    rule = db.query(TenantRuleModel).filter(TenantRuleModel.tenant_id == current_user.tenant_id).first()
    if rule:
        rule.max_temperature = max_temp
        rule.max_current = max_curr
    else:
        db.add(TenantRuleModel(tenant_id=current_user.tenant_id, max_temperature=max_temp, max_current=max_curr))
        
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CONFIGURAR_REGRAS", details=f"Limites da organização atualizados: temperatura={max_temp}°C, corrente={max_curr}A."))
    db.commit()
    return {"status": "success", "message": f"Regras de alarme atualizadas para Limite Térmico: {max_temp}°C"}

@app.post("/api/panels/add")
def add_panel(data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in ["CEO", "Gerente"]:
        raise HTTPException(status_code=403, detail="Acesso restrito para adicionar painéis.")
    
    raw_id = data.get("id", "")
    raw_name = data.get("name", "")
    raw_location = data.get("location", "")
    raw_voltage = data.get("voltage", "380V AC")
    raw_component = data.get("component", "Disjuntor Principal")
    raw_qr = data.get("gateway_qr", f"QR-GW-{random.randint(2000,9999)}")
    if not all(isinstance(value, str) for value in (raw_id, raw_name, raw_location, raw_voltage, raw_component, raw_qr)):
        raise HTTPException(status_code=422, detail="Os campos do painel devem ser textos.")
    panel_id = raw_id.strip().upper()
    name = raw_name.strip()
    location = raw_location.strip()
    voltage = raw_voltage.strip()
    component = raw_component.strip()
    gateway_qr = raw_qr.strip()

    if not panel_id or not name or not location:
        raise HTTPException(status_code=400, detail="Preencha os campos obrigatórios do painel.")
    import re
    if not re.fullmatch(r"[A-Z0-9_-]{1,100}", panel_id):
        raise HTTPException(status_code=422, detail="ID deve conter apenas letras, números, hífen e sublinhado.")

    if len(panel_id) > 100 or len(name) > 200 or len(location) > 200 or len(voltage) > 80 or len(component) > 160 or len(gateway_qr) > 160:
        raise HTTPException(status_code=422, detail="Um ou mais campos excedem o tamanho permitido.")
    existing = db.query(PanelModel).filter(PanelModel.id == panel_id).first()
    if existing:
        raise HTTPException(status_code=409, detail="Este ID de painel já está em uso. IDs de painéis são globalmente únicos nesta versão.")
    subscription = db.query(TenantSubscriptionModel).filter(TenantSubscriptionModel.tenant_id == current_user.tenant_id).first()
    panel_count = db.query(PanelModel).filter(PanelModel.tenant_id == current_user.tenant_id).count()
    if subscription and panel_count >= subscription.max_panels:
        raise HTTPException(status_code=409, detail="Limite de painéis do plano atingido.")

    new_panel = PanelModel(
        id=panel_id, tenant_id=current_user.tenant_id, name=name, location=location,
        voltage=voltage, current=round(random.uniform(100.0, 400.0), 2),
        temperature=round(random.uniform(45.0, 68.0), 2), gateway_qr=gateway_qr,
        status="NORMAL", component=component, predictive_failure_hours=180,
        recommendation="Painel recém-comissionado. Operação estável.", esg_waste_kwh=1.0
    )
    db.add(new_panel)
    db.add(PanelHistoryModel(panel_id=panel_id, tenant_id=current_user.tenant_id, timestamp=datetime.utcnow().isoformat(timespec="seconds") + "Z", temperature=new_panel.temperature, current=new_panel.current, voltage=new_panel.voltage))
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="ADICIONAR_PAINEL", details=f"Painel {panel_id} adicionado."))
    db.commit()
    return {"status": "success", "message": f"Painel {panel_id} adicionado com sucesso!"}

def get_effective_alarm_rule(db: Session, tenant_id: str, panel_id: str):
    panel_rule = db.query(PanelAlarmRuleModel).filter(PanelAlarmRuleModel.tenant_id == tenant_id, PanelAlarmRuleModel.panel_id == panel_id).first()
    if panel_rule:
        return panel_rule
    tenant_rule = db.query(TenantRuleModel).filter(TenantRuleModel.tenant_id == tenant_id).first()
    return type("EffectiveRule", (), {"max_temperature": tenant_rule.max_temperature if tenant_rule else 85.0,
        "max_current": tenant_rule.max_current if tenant_rule else 500.0,
        "temperature_hysteresis": 3.0, "current_hysteresis": 10.0})()


def condition_is_active(value: float, limit: float, hysteresis: float, alarm_exists: bool) -> bool:
    # An open alarm remains active until the reading falls below limit - hysteresis.
    return value > (limit - hysteresis if alarm_exists else limit)


@app.get("/api/panels/{panel_id}/alarm-rules", tags=["Alarm Rules"])
def get_panel_alarm_rules(panel_id: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    panel = db.query(PanelModel).filter(PanelModel.id == panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    custom = db.query(PanelAlarmRuleModel).filter(PanelAlarmRuleModel.tenant_id == current_user.tenant_id, PanelAlarmRuleModel.panel_id == panel_id).first()
    rule = get_effective_alarm_rule(db, current_user.tenant_id, panel_id)
    return {"panel_id": panel_id, "custom": custom is not None, "max_temperature": rule.max_temperature,
            "max_current": rule.max_current, "temperature_hysteresis": rule.temperature_hysteresis,
            "current_hysteresis": rule.current_hysteresis}


@app.put("/api/panels/{panel_id}/alarm-rules", tags=["Alarm Rules"])
def save_panel_alarm_rules(panel_id: str, data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}:
        raise HTTPException(status_code=403, detail="Sem permissão para alterar regras de alarme.")
    panel = db.query(PanelModel).filter(PanelModel.id == panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    import math
    try:
        values = {"max_temperature": float(data.get("max_temperature", 85)), "max_current": float(data.get("max_current", 500)),
                  "temperature_hysteresis": float(data.get("temperature_hysteresis", 3)), "current_hysteresis": float(data.get("current_hysteresis", 10))}
    except (TypeError, ValueError):
        raise HTTPException(status_code=422, detail="Todos os limites devem ser numéricos.")
    if not all(math.isfinite(v) for v in values.values()):
        raise HTTPException(status_code=422, detail="Valores não finitos não são permitidos.")
    if not 0 < values["max_temperature"] <= 250 or not 0 < values["max_current"] <= 100000:
        raise HTTPException(status_code=422, detail="Limites máximos fora dos intervalos permitidos.")
    if not 0 <= values["temperature_hysteresis"] < values["max_temperature"] or not 0 <= values["current_hysteresis"] < values["max_current"]:
        raise HTTPException(status_code=422, detail="A histerese deve ser não negativa e menor que o limite correspondente.")
    rule = db.query(PanelAlarmRuleModel).filter(PanelAlarmRuleModel.tenant_id == current_user.tenant_id, PanelAlarmRuleModel.panel_id == panel_id).first()
    if rule is None:
        rule = PanelAlarmRuleModel(tenant_id=current_user.tenant_id, panel_id=panel_id)
        db.add(rule)
    for key, value in values.items(): setattr(rule, key, value)
    rule.updated_at = datetime.utcnow()
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role,
        action="CONFIGURAR_REGRAS_PAINEL", details=f"Regras atualizadas para {panel_id}: temperatura={values['max_temperature']}°C, corrente={values['max_current']}A."))
    db.commit()
    return {"status": "success", "message": "Regras específicas do painel atualizadas.", "panel_id": panel_id, **values}


@app.delete("/api/panels/{panel_id}/alarm-rules", tags=["Alarm Rules"])
def delete_panel_alarm_rules(panel_id: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}:
        raise HTTPException(status_code=403, detail="Sem permissão para alterar regras de alarme.")
    panel = db.query(PanelModel).filter(PanelModel.id == panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    rule = db.query(PanelAlarmRuleModel).filter(PanelAlarmRuleModel.tenant_id == current_user.tenant_id, PanelAlarmRuleModel.panel_id == panel_id).first()
    if rule:
        db.delete(rule)
        db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role,
            action="RESTAURAR_REGRAS_PAINEL", details=f"Regras específicas removidas do painel {panel_id}; aplicados os padrões da organização."))
        db.commit()
    return {"status": "success", "custom": False, "panel_id": panel_id, "message": "O painel está usando os limites padrão da organização."}


@app.post("/api/telemetry", tags=["Industrial Ingestion"])
def ingest_telemetry(data: TelemetryData, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    import math
    if not data.panel_id or len(data.panel_id) > 100:
        raise HTTPException(status_code=422, detail="ID de painel inválido.")
    if not math.isfinite(data.temperature) or not -50 <= data.temperature <= 250:
        raise HTTPException(status_code=422, detail="Temperatura fora do intervalo permitido (-50 a 250 °C).")
    if not math.isfinite(data.current) or not 0 <= data.current <= 100000:
        raise HTTPException(status_code=422, detail="Corrente fora do intervalo permitido (0 a 100000 A).")
    if data.voltage is not None and len(data.voltage) > 80:
        raise HTTPException(status_code=422, detail="Tensão demasiado longa.")
    if data.gateway_qr is not None and len(data.gateway_qr) > 160:
        raise HTTPException(status_code=422, detail="Identificador do gateway demasiado longo.")
    panel = db.query(PanelModel).filter(PanelModel.id == data.panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail=f"Painel {data.panel_id} não encontrado.")
    
    rule = get_effective_alarm_rule(db, current_user.tenant_id, panel.id)
    temp_limit, current_limit = rule.max_temperature, rule.max_current
    
    panel.temperature = data.temperature
    panel.current = data.current
    if data.voltage:
        panel.voltage = data.voltage
        
    if data.temperature > temp_limit or data.current > current_limit:
        panel.status = "CRITICO"
        reasons = []
        if data.temperature > temp_limit:
            reasons.append(f"temperatura acima de {temp_limit}°C")
        if data.current > current_limit:
            reasons.append(f"corrente acima de {current_limit} A")
        panel.recommendation = "ALERTA: " + " e ".join(reasons) + "."
    else:
        panel.status = "NORMAL"
        panel.recommendation = "Sistema a operar dentro dos parâmetros configurados."
        
    # Alarm lifecycle: one open alarm per panel/type; repeated events update the same record.
    open_alarms = db.query(AlarmModel).filter(
        AlarmModel.tenant_id == current_user.tenant_id,
        AlarmModel.panel_id == panel.id,
        AlarmModel.status.in_(["ACTIVE", "ACKNOWLEDGED"]),
    ).all()
    open_types = {item.alarm_type for item in open_alarms}
    active_conditions = []
    if condition_is_active(data.temperature, temp_limit, rule.temperature_hysteresis, "HIGH_TEMPERATURE" in open_types):
        active_conditions.append(("HIGH_TEMPERATURE", "CRITICAL", f"Temperatura {data.temperature} °C acima do limite configurado de {temp_limit} °C."))
    if condition_is_active(data.current, current_limit, rule.current_hysteresis, "HIGH_CURRENT" in open_types):
        active_conditions.append(("HIGH_CURRENT", "CRITICAL", f"Corrente {data.current} A acima do limite configurado de {current_limit} A."))
    active_types = {condition[0] for condition in active_conditions}
    now = datetime.utcnow()
    for alarm_type, severity, message in active_conditions:
        alarm = next((item for item in open_alarms if item.alarm_type == alarm_type), None)
        if alarm:
            alarm.last_triggered_at = now
            alarm.message = message
            alarm.severity = severity
        else:
            db.add(AlarmModel(tenant_id=current_user.tenant_id, panel_id=panel.id, alarm_type=alarm_type, severity=severity, message=message, status="ACTIVE", first_triggered_at=now, last_triggered_at=now))
    for alarm in open_alarms:
        if alarm.alarm_type in {"HIGH_TEMPERATURE", "HIGH_CURRENT"} and alarm.alarm_type not in active_types:
            alarm.status = "CLEARED"
            alarm.cleared_at = now
    db.add(PanelHistoryModel(panel_id=panel.id, tenant_id=current_user.tenant_id, timestamp=now.isoformat(timespec="seconds") + "Z", temperature=panel.temperature, current=panel.current, voltage=panel.voltage))
    db.commit()
    return {"status": "success", "message": f"Telemetria de {data.panel_id} atualizada com base nos limites configurados.", "active_alarms": len(active_conditions)}

@app.post("/api/demo/force-failure")
async def force_failure(data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in ["CEO", "Gerente"]:
        raise HTTPException(status_code=403, detail="Apenas CEO ou Gerente pode executar simulações.")
    raw_panel_id = data.get("panel_id", "")
    panel_id = raw_panel_id.strip() if isinstance(raw_panel_id, str) else ""
    panel = db.query(PanelModel).filter(PanelModel.id == panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel não encontrado.")
    
    # Simulação explícita: atualiza telemetria e cria/atualiza alarmes persistentes.
    # Não depende de PLC, Modbus ou equipamento físico.
    rule = db.query(TenantRuleModel).filter(TenantRuleModel.tenant_id == current_user.tenant_id).first()
    temp_limit = rule.max_temperature if rule else 85.0
    current_limit = rule.max_current if rule else 500.0
    panel.status = "CRITICO"
    panel.temperature = max(94.5, temp_limit + 5.0)
    panel.current = max(float(panel.current or 0), current_limit + 25.0)
    panel.predictive_failure_hours = 2
    panel.recommendation = "ALERTA CRÍTICO: falha elétrica simulada para teste da Central de Alarmes."

    now = datetime.utcnow()
    simulated_conditions = [
        ("HIGH_TEMPERATURE", f"Temperatura simulada {panel.temperature} °C acima do limite configurado de {temp_limit} °C."),
        ("HIGH_CURRENT", f"Corrente simulada {panel.current} A acima do limite configurado de {current_limit} A."),
    ]
    open_alarms = db.query(AlarmModel).filter(
        AlarmModel.tenant_id == current_user.tenant_id,
        AlarmModel.panel_id == panel.id,
        AlarmModel.status.in_(["ACTIVE", "ACKNOWLEDGED"]),
    ).all()
    for alarm_type, message in simulated_conditions:
        existing_alarm = next((a for a in open_alarms if a.alarm_type == alarm_type), None)
        if existing_alarm:
            existing_alarm.last_triggered_at = now
            existing_alarm.message = message
            existing_alarm.severity = "CRITICAL"
            # Preserve ACKNOWLEDGED status when the operator already acknowledged it.
        else:
            db.add(AlarmModel(
                tenant_id=current_user.tenant_id, panel_id=panel.id,
                alarm_type=alarm_type, severity="CRITICAL", message=message,
                status="ACTIVE", first_triggered_at=now, last_triggered_at=now,
            ))

    db.add(PanelHistoryModel(panel_id=panel.id, tenant_id=current_user.tenant_id, timestamp=now.isoformat(timespec="seconds") + "Z", temperature=panel.temperature, current=panel.current, voltage=panel.voltage))
    db.add(AuditLogModel(
        tenant_id=current_user.tenant_id,
        user_name=current_user.username,
        role=current_user.role,
        action="FORCAR_FALHA",
        details=f"Falha crítica simulada no painel {panel_id}."
    ))
    db.commit()

    webhook_delivery = "NOT_CONFIGURED"
    webhook_http_status = None
    sub = db.query(TenantSubscriptionModel).filter(TenantSubscriptionModel.tenant_id == current_user.tenant_id).first()
    if sub and sub.webhook_url:
        try:
            safe_url = validate_webhook_url(sub.webhook_url)
            async with httpx.AsyncClient(follow_redirects=False, timeout=5.0) as client:
                webhook_response = await client.post(safe_url, json={"event": "CRITICAL_ALERT", "panel_id": panel.id, "temperature": panel.temperature, "tenant_id": current_user.tenant_id})
                webhook_http_status = webhook_response.status_code
                if not 200 <= webhook_response.status_code < 300:
                    raise RuntimeError("Webhook retornou um status HTTP sem confirmação de sucesso.")
            webhook_delivery = "DELIVERED"
        except Exception:
            # Uma falha externa não desfaz o alarme; o retorno informa que a entrega não foi confirmada.
            webhook_delivery = "FAILED"

    await manager.broadcast({"event": "PANEL_UPDATE", "panel_id": panel.id, "status": "CRITICO"}, current_user.tenant_id)
    active_alarm_count = db.query(AlarmModel).filter(
        AlarmModel.tenant_id == current_user.tenant_id,
        AlarmModel.panel_id == panel.id,
        AlarmModel.status.in_(["ACTIVE", "ACKNOWLEDGED"]),
    ).count()
    return {"status": "success", "message": f"Falha simulada em {panel.id}. {active_alarm_count} alarme(s) ativo(s) ou reconhecido(s) registrado(s).", "active_alarms": active_alarm_count, "webhook_delivery": webhook_delivery, "webhook_http_status": webhook_http_status}

@app.get("/api/panels")
def get_panels(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    try:
        panels = db.query(PanelModel).filter(PanelModel.tenant_id == current_user.tenant_id).all()
        
        panels_data = []
        offline_after_seconds = int(os.getenv("DEVICE_OFFLINE_AFTER_SECONDS", "300"))
        offline_after_seconds = max(30, min(offline_after_seconds, 86400))
        for p in panels:
            curr = p.current
            temp = p.temperature
            if os.getenv("SIMULATE_TELEMETRY", "0").lower() in {"1", "true", "yes"} and p.status != "CRITICO":
                curr = round(max(0.0, p.current + random.uniform(-1.2, 1.2)), 2)
                temp = round(max(-50.0, p.temperature + random.uniform(-0.3, 0.3)), 2)

            # Connectivity is based on real device-token activity, not seeded/demo history.
            # A panel without an active device credential is reported as NOT_CONFIGURED.
            device_rows = db.query(DeviceCredentialModel).filter(
                DeviceCredentialModel.tenant_id == current_user.tenant_id,
                DeviceCredentialModel.panel_id == p.id,
                DeviceCredentialModel.revoked_at.is_(None),
            ).all()
            last_seen = max((x.last_used_at for x in device_rows if x.last_used_at), default=None)
            if not device_rows:
                connectivity_status = "NOT_CONFIGURED"
            elif last_seen is None:
                connectivity_status = "NEVER_CONNECTED"
            else:
                age_seconds = max(0, (datetime.utcnow() - last_seen).total_seconds())
                connectivity_status = "ONLINE" if age_seconds <= offline_after_seconds else "OFFLINE"

            panels_data.append({
                "id": p.id, "name": p.name, "location": p.location, "voltage": p.voltage,
                "voltage_v": p.voltage_v, "active_power_kw": p.active_power_kw,
                "active_energy_kwh": p.active_energy_kwh, "frequency_hz": p.frequency_hz,
                "power_factor": p.power_factor,
                "current": curr, "temperature": temp, "temperature_measured": p.temperature_measured, "gateway_qr": p.gateway_qr,
                "status": p.status, "component": p.component, "predictive_failure_hours": p.predictive_failure_hours,
                "recommendation": p.recommendation, "esg_waste_kwh": p.esg_waste_kwh, "assigned_to": p.assigned_to,
                "connectivity_status": connectivity_status,
                "last_seen_at": last_seen.isoformat() + "Z" if last_seen else None,
                "offline_after_seconds": offline_after_seconds,
            })

        logs = db.query(AuditLogModel).filter(AuditLogModel.tenant_id == current_user.tenant_id).order_by(AuditLogModel.timestamp.desc()).limit(50).all()
        sub = db.query(TenantSubscriptionModel).filter(TenantSubscriptionModel.tenant_id == current_user.tenant_id).first()
        rule = db.query(TenantRuleModel).filter(TenantRuleModel.tenant_id == current_user.tenant_id).first()
        work_orders = db.query(WorkOrderModel).filter(WorkOrderModel.tenant_id == current_user.tenant_id).all()
        work_order_ids = [w.id for w in work_orders]
        work_order_history = []
        if work_order_ids:
            work_order_history = db.query(WorkOrderHistoryModel).filter(
                WorkOrderHistoryModel.tenant_id == current_user.tenant_id,
                WorkOrderHistoryModel.work_order_id.in_(work_order_ids),
            ).order_by(WorkOrderHistoryModel.timestamp.desc(), WorkOrderHistoryModel.id.desc()).limit(1000).all()
        history_by_order = {}
        for event in work_order_history:
            history_by_order.setdefault(event.work_order_id, []).append({
                "timestamp": utc_iso(event.timestamp), "user_name": event.user_name,
                "role": event.role, "action": event.action, "details": event.details,
                "resolution": event.resolution,
            })
        for events in history_by_order.values():
            events.reverse()
        checklist_by_order = {}
        if work_order_ids:
            checklist_items = db.query(WorkOrderChecklistItemModel).filter(
                WorkOrderChecklistItemModel.tenant_id == current_user.tenant_id,
                WorkOrderChecklistItemModel.work_order_id.in_(work_order_ids),
            ).order_by(WorkOrderChecklistItemModel.id.asc()).all()
            for item in checklist_items:
                checklist_by_order.setdefault(item.work_order_id, []).append({
                    "id": item.id, "label": item.label, "is_completed": item.is_completed,
                    "completed_at": utc_iso(item.completed_at), "completed_by": item.completed_by,
                })
        inventory = db.query(InventoryModel).filter(InventoryModel.tenant_id == current_user.tenant_id).all()
        alarms = db.query(AlarmModel).filter(AlarmModel.tenant_id == current_user.tenant_id).order_by(AlarmModel.first_triggered_at.desc()).limit(200).all()
        
        history_data = {}
        for p in panels:
            # Bound dashboard payload size; dedicated endpoint exposes paginated history.
            h_records = db.query(PanelHistoryModel).filter(
                PanelHistoryModel.panel_id == p.id,
                PanelHistoryModel.tenant_id == current_user.tenant_id,
            ).order_by(PanelHistoryModel.id.desc()).limit(500).all()
            h_records.reverse()
            history_data[p.id] = [{
                "timestamp": h.timestamp, "temperature": h.temperature,
                "temperature_measured": h.temperature_measured,
                "current": h.current, "voltage": h.voltage, "voltage_v": h.voltage_v,
                "active_power_kw": h.active_power_kw, "active_energy_kwh": h.active_energy_kwh,
                "frequency_hz": h.frequency_hz, "power_factor": h.power_factor,
            } for h in h_records]

        logs_data = [{"timestamp": str(l.timestamp), "user_name": l.user_name, "role": l.role, "action": l.action, "details": l.details} for l in logs]
        wo_data = [{"id": w.id, "alarm_id": w.alarm_id, "panel_id": w.panel_id, "title": w.title, "priority": w.priority, "status": w.status, "assigned_technician": w.assigned_technician, "created_at": utc_iso(w.created_at), "due_at": utc_iso(w.due_at), "started_at": utc_iso(w.started_at), "completed_at": utc_iso(w.completed_at), "description": w.description, "resolution": w.resolution, "history": history_by_order.get(w.id, []), "checklist": checklist_by_order.get(w.id, [])} for w in work_orders]
        inv_data = [{"id": i.id, "part_name": i.part_name, "category": i.category, "quantity_in_stock": i.quantity_in_stock, "min_threshold": i.min_threshold, "unit_cost": i.unit_cost} for i in inventory]
        
        billing_info = {
            "plan_name": sub.plan_name if sub else "Enterprise",
            "status": sub.status if sub else "ACTIVE",
            "expires_at": str(sub.expires_at) if sub and sub.expires_at else None,
            "max_panels": sub.max_panels if sub else 25,
            "webhook_url": sub.webhook_url if sub else "",
            "max_temperature": rule.max_temperature if rule else 85.0,
            "max_current": rule.max_current if rule else 500.0
        }

        work_order_by_alarm = {w.alarm_id: w.id for w in work_orders if w.alarm_id is not None}
        alarm_data = [{"id": a.id, "work_order_id": work_order_by_alarm.get(a.id), "panel_id": a.panel_id, "alarm_type": a.alarm_type, "severity": a.severity,
                       "message": a.message, "status": a.status, "first_triggered_at": a.first_triggered_at.isoformat() + "Z",
                       "last_triggered_at": a.last_triggered_at.isoformat() + "Z", "acknowledged_at": a.acknowledged_at.isoformat() + "Z" if a.acknowledged_at else None,
                       "acknowledged_by": a.acknowledged_by, "cleared_at": a.cleared_at.isoformat() + "Z" if a.cleared_at else None} for a in alarms]
        return {
            "panels": panels_data, "audit_logs": logs_data, "work_orders": wo_data,
            "inventory": inv_data, "history": history_data, "alarms": alarm_data, "tenant_id": current_user.tenant_id, "billing": billing_info
        }
    except Exception:
        raise HTTPException(status_code=500, detail="Não foi possível carregar os dados do painel.")

@app.get("/api/audit-logs")
def list_audit_logs(
    limit: int = 100,
    offset: int = 0,
    search: Optional[str] = None,
    user_name: Optional[str] = None,
    action: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role != "CEO":
        raise HTTPException(status_code=403, detail="Somente CEO pode consultar o registro de auditoria.")
    if limit < 1 or limit > 200 or offset < 0:
        raise HTTPException(status_code=422, detail="Paginação inválida. O limite deve ser de 1 a 200 e o deslocamento não pode ser negativo.")
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=422, detail="A data inicial deve ser anterior ou igual à data final.")

    query = db.query(AuditLogModel).filter(AuditLogModel.tenant_id == current_user.tenant_id)
    term = (search or "").strip()
    actor = (user_name or "").strip()
    action_name = (action or "").strip()
    if len(term) > 200 or len(actor) > 120 or len(action_name) > 100:
        raise HTTPException(status_code=422, detail="Um dos filtros excede o tamanho permitido.")
    if term:
        pattern = f"%{term}%"
        query = query.filter(or_(
            AuditLogModel.user_name.ilike(pattern), AuditLogModel.role.ilike(pattern),
            AuditLogModel.action.ilike(pattern), AuditLogModel.details.ilike(pattern),
        ))
    if actor:
        query = query.filter(AuditLogModel.user_name.ilike(f"%{actor}%"))
    if action_name:
        query = query.filter(AuditLogModel.action == action_name)
    if date_from:
        query = query.filter(AuditLogModel.timestamp >= datetime.combine(date_from, datetime.min.time()))
    if date_to:
        query = query.filter(AuditLogModel.timestamp < datetime.combine(date_to + timedelta(days=1), datetime.min.time()))

    total = query.count()
    rows = query.order_by(AuditLogModel.timestamp.desc(), AuditLogModel.id.desc()).offset(offset).limit(limit).all()
    actions = [row[0] for row in db.query(AuditLogModel.action).filter(
        AuditLogModel.tenant_id == current_user.tenant_id
    ).distinct().order_by(AuditLogModel.action.asc()).all()]
    return {
        "items": [{"id": row.id, "timestamp": utc_iso(row.timestamp), "user_name": row.user_name,
                   "role": row.role, "action": row.action, "details": row.details} for row in rows],
        "total": total, "limit": limit, "offset": offset, "has_more": offset + len(rows) < total,
        "actions": actions,
    }

@app.get("/api/telemetry/history/{panel_id}", tags=["Industrial Ingestion"])
def get_telemetry_history(
    panel_id: str,
    limit: int = 100,
    before_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Retorna histórico paginado, sempre limitado ao tenant autenticado."""
    if limit < 1 or limit > 1000:
        raise HTTPException(status_code=422, detail="O limite deve estar entre 1 e 1000 registros.")
    panel = db.query(PanelModel).filter(
        PanelModel.id == panel_id,
        PanelModel.tenant_id == current_user.tenant_id,
    ).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    query = db.query(PanelHistoryModel).filter(
        PanelHistoryModel.panel_id == panel_id,
        PanelHistoryModel.tenant_id == current_user.tenant_id,
    )
    if before_id is not None:
        query = query.filter(PanelHistoryModel.id < before_id)
    rows = query.order_by(PanelHistoryModel.id.desc()).limit(limit + 1).all()
    has_more = len(rows) > limit
    rows = rows[:limit]
    next_before_id = rows[-1].id if has_more and rows else None
    rows.reverse()
    return {
        "panel_id": panel_id,
        "count": len(rows),
        "limit": limit,
        "has_more": has_more,
        "next_before_id": next_before_id,
        "history": [
            {"id": row.id, "timestamp": row.timestamp, "temperature": row.temperature,
             "temperature_measured": row.temperature_measured,
             "current": row.current, "voltage": row.voltage, "voltage_v": row.voltage_v,
             "active_power_kw": row.active_power_kw, "active_energy_kwh": row.active_energy_kwh,
             "frequency_hz": row.frequency_hz, "power_factor": row.power_factor}
            for row in rows
        ],
    }


@app.get("/api/alarms")
def list_alarms(status_filter: Optional[str] = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    query = db.query(AlarmModel).filter(AlarmModel.tenant_id == current_user.tenant_id)
    if status_filter:
        if status_filter not in {"ACTIVE", "ACKNOWLEDGED", "CLEARED"}:
            raise HTTPException(status_code=422, detail="Filtro de estado de alarme inválido.")
        query = query.filter(AlarmModel.status == status_filter)
    alarms = query.order_by(AlarmModel.first_triggered_at.desc()).limit(500).all()
    linked_orders = dict(db.query(WorkOrderModel.alarm_id, WorkOrderModel.id).filter(
        WorkOrderModel.tenant_id == current_user.tenant_id, WorkOrderModel.alarm_id.is_not(None)
    ).all())
    return [{"id": a.id, "work_order_id": linked_orders.get(a.id), "panel_id": a.panel_id, "alarm_type": a.alarm_type, "severity": a.severity,
             "message": a.message, "status": a.status, "first_triggered_at": a.first_triggered_at.isoformat() + "Z",
             "last_triggered_at": a.last_triggered_at.isoformat() + "Z", "acknowledged_at": a.acknowledged_at.isoformat() + "Z" if a.acknowledged_at else None,
             "acknowledged_by": a.acknowledged_by, "cleared_at": a.cleared_at.isoformat() + "Z" if a.cleared_at else None}
            for a in alarms]


@app.post("/api/alarms/{alarm_id}/acknowledge")
def acknowledge_alarm(alarm_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para reconhecer alarmes.")
    alarm = db.query(AlarmModel).filter(AlarmModel.id == alarm_id, AlarmModel.tenant_id == current_user.tenant_id).first()
    if not alarm:
        raise HTTPException(status_code=404, detail="Alarme não encontrado nesta organização.")
    if alarm.status == "CLEARED":
        raise HTTPException(status_code=409, detail="Um alarme encerrado não pode ser reconhecido.")
    if alarm.status != "ACKNOWLEDGED":
        alarm.status = "ACKNOWLEDGED"
        alarm.acknowledged_at = datetime.utcnow()
        alarm.acknowledged_by = current_user.username
        db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="RECONHECER_ALARME", details=f"Alarme {alarm_id} reconhecido."))
        db.commit()
    return {"status": "success", "id": alarm.id, "alarm_status": alarm.status}


@app.post("/api/alarms/{alarm_id}/work-order", status_code=201, tags=["Work Orders"])
def create_work_order_from_alarm(alarm_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Create one traceable maintenance order from an open alarm."""
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para criar ordens de serviço.")
    alarm = db.query(AlarmModel).filter(
        AlarmModel.id == alarm_id, AlarmModel.tenant_id == current_user.tenant_id
    ).first()
    if not alarm:
        raise HTTPException(status_code=404, detail="Alarme não encontrado nesta organização.")
    existing = db.query(WorkOrderModel).filter(
        WorkOrderModel.tenant_id == current_user.tenant_id, WorkOrderModel.alarm_id == alarm.id
    ).first()
    if existing:
        return JSONResponse(status_code=200, content={
            "status": "already_exists", "id": existing.id, "alarm_id": alarm.id,
            "panel_id": existing.panel_id, "title": existing.title, "priority": existing.priority,
            "order_status": existing.status, "due_at": utc_iso(existing.due_at),
            "started_at": utc_iso(existing.started_at), "completed_at": utc_iso(existing.completed_at),
            "resolution": existing.resolution,
        })
    if alarm.status == "CLEARED":
        raise HTTPException(status_code=409, detail="Não é possível abrir uma ordem a partir de um alarme encerrado.")
    panel = db.query(PanelModel).filter(
        PanelModel.id == alarm.panel_id, PanelModel.tenant_id == current_user.tenant_id
    ).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel associado não encontrado nesta organização.")
    title = f"Investigar {alarm.alarm_type.replace('_', ' ').lower()} — {panel.name}"
    order = WorkOrderModel(
        tenant_id=current_user.tenant_id, alarm_id=alarm.id, panel_id=panel.id,
        title=title[:200], priority="CRITICA" if alarm.severity == "CRITICAL" else "ALTA",
        status="ABERTA", description=(
            f"Criada a partir do alarme #{alarm.id}. {alarm.message}\n"
            f"Primeira ocorrência: {alarm.first_triggered_at.isoformat()}Z; "
            f"última ocorrência: {alarm.last_triggered_at.isoformat()}Z."
        )[:4000],
    )
    db.add(order)
    db.add(AuditLogModel(
        tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role,
        action="CRIAR_ORDEM_A_PARTIR_DE_ALARME",
        details=f"Ordem de serviço criada para o alarme {alarm.id} no painel {panel.id}.",
    ))
    try:
        db.flush()
        checklist_labels = [
            "Inspecionar visualmente o painel e a área",
            "Verificar a condição indicada pelo alarme",
            "Registrar a solução e validar o retorno à condição normal",
        ]
        db.add_all([
            WorkOrderChecklistItemModel(
                tenant_id=current_user.tenant_id, work_order_id=order.id, label=label
            ) for label in checklist_labels
        ])
        db.add(WorkOrderHistoryModel(
            tenant_id=current_user.tenant_id, work_order_id=order.id,
            user_name=current_user.username, role=current_user.role, action="CRIADA",
            details=f"Ordem criada a partir do alarme #{alarm.id} no painel {panel.id}. Título: {order.title}. Prioridade: {order.priority}. Estado inicial: ABERTA.",
        ))
        db.commit()
    except Exception:
        db.rollback()
        existing = db.query(WorkOrderModel).filter(
            WorkOrderModel.tenant_id == current_user.tenant_id, WorkOrderModel.alarm_id == alarm.id
        ).first()
        if existing:
            return JSONResponse(status_code=200, content={
                "status": "already_exists", "id": existing.id, "alarm_id": alarm.id,
                "panel_id": existing.panel_id, "title": existing.title, "priority": existing.priority,
                "order_status": existing.status, "due_at": utc_iso(existing.due_at),
                "started_at": utc_iso(existing.started_at), "completed_at": utc_iso(existing.completed_at),
                "resolution": existing.resolution,
            })
        raise
    db.refresh(order)
    return {"status": "created", "id": order.id, "alarm_id": alarm.id,
            "panel_id": order.panel_id, "title": order.title, "priority": order.priority,
            "order_status": order.status, "due_at": utc_iso(order.due_at),
            "started_at": utc_iso(order.started_at), "completed_at": utc_iso(order.completed_at),
            "resolution": order.resolution}


@app.post("/api/alarms/{alarm_id}/clear")
def clear_alarm(alarm_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para encerrar alarmes.")
    alarm = db.query(AlarmModel).filter(AlarmModel.id == alarm_id, AlarmModel.tenant_id == current_user.tenant_id).first()
    if not alarm:
        raise HTTPException(status_code=404, detail="Alarme não encontrado nesta organização.")
    if alarm.status != "CLEARED":
        alarm.status = "CLEARED"
        alarm.cleared_at = datetime.utcnow()
        db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="ENCERRAR_ALARME", details=f"Alarme {alarm_id} encerrado manualmente."))
        db.commit()
    return {"status": "success", "id": alarm.id, "alarm_status": alarm.status}


@app.get("/api/panels/qr/{qr_code}")
def get_panel_by_qr(qr_code: str, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    panel = db.query(PanelModel).filter(PanelModel.gateway_qr == qr_code, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        panel = db.query(PanelModel).filter(PanelModel.id == qr_code, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Nenhum cubículo encontrado para este código QR ou ID.")
    device_rows = db.query(DeviceCredentialModel).filter(
        DeviceCredentialModel.tenant_id == current_user.tenant_id,
        DeviceCredentialModel.panel_id == panel.id,
        DeviceCredentialModel.revoked_at.is_(None),
    ).all()
    last_seen = max((item.last_used_at for item in device_rows if item.last_used_at), default=None)
    offline_after_seconds = max(30, min(int(os.getenv("DEVICE_OFFLINE_AFTER_SECONDS", "300")), 86400))
    if not device_rows:
        connectivity_status = "NOT_CONFIGURED"
    elif last_seen is None:
        connectivity_status = "NEVER_CONNECTED"
    else:
        age_seconds = max(0, (datetime.utcnow() - last_seen).total_seconds())
        connectivity_status = "ONLINE" if age_seconds <= offline_after_seconds else "OFFLINE"
    return {
        "id": panel.id, "name": panel.name, "location": panel.location, "voltage": panel.voltage,
        "voltage_v": panel.voltage_v, "active_power_kw": panel.active_power_kw,
        "active_energy_kwh": panel.active_energy_kwh, "frequency_hz": panel.frequency_hz,
        "power_factor": panel.power_factor,
        "current": panel.current, "temperature": panel.temperature, "gateway_qr": panel.gateway_qr,
        "temperature_measured": panel.temperature_measured,
        "status": panel.status, "component": panel.component, "predictive_failure_hours": panel.predictive_failure_hours,
        "recommendation": panel.recommendation, "esg_waste_kwh": panel.esg_waste_kwh, "assigned_to": panel.assigned_to,
        "connectivity_status": connectivity_status, "last_seen_at": last_seen.isoformat() + "Z" if last_seen else None,
    }

@app.post("/api/inventory/update")
def update_inventory_item(data: dict, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in ["CEO", "Gerente"]:
        raise HTTPException(status_code=403, detail="Acesso restrito.")
    
    part_name = data.get("part_name")
    category = data.get("category", "Elétrica")
    raw_qty = data.get("quantity_change", 0)
    raw_min = data.get("min_threshold", 3)
    try:
        qty_change = int(raw_qty)
        unit_cost = float(data.get("unit_cost", 150.0))
        min_threshold = int(raw_min)
    except (TypeError, ValueError, OverflowError):
        raise HTTPException(status_code=422, detail="Quantidade, custo e estoque mínimo devem ser numéricos.")
    if isinstance(raw_qty, float) and not raw_qty.is_integer():
        raise HTTPException(status_code=422, detail="A quantidade deve ser um número inteiro.")
    if isinstance(raw_min, float) and not raw_min.is_integer():
        raise HTTPException(status_code=422, detail="O estoque mínimo deve ser um número inteiro.")
    import math
    if not isinstance(part_name, str) or not part_name.strip() or len(part_name.strip()) > 160:
        raise HTTPException(status_code=422, detail="Nome da peça inválido.")
    if not isinstance(category, str) or len(category.strip()) > 120:
        raise HTTPException(status_code=422, detail="Categoria inválida.")
    part_name = part_name.strip()
    if not math.isfinite(unit_cost):
        raise HTTPException(status_code=422, detail="Custo unitário inválido.")
    if qty_change < 0 and not db.query(InventoryModel).filter(InventoryModel.part_name == part_name, InventoryModel.tenant_id == current_user.tenant_id).first():
        raise HTTPException(status_code=422, detail="Não é possível criar um item com ajuste negativo.")
    if qty_change < -100000 or qty_change > 100000 or min_threshold < 0 or not (0 <= unit_cost <= 100000000):
        raise HTTPException(status_code=422, detail="Quantidade, estoque mínimo ou custo fora dos limites.")

    item = db.query(InventoryModel).filter(InventoryModel.part_name == part_name, InventoryModel.tenant_id == current_user.tenant_id).first()
    if item:
        new_quantity = item.quantity_in_stock + qty_change
        if new_quantity < 0:
            raise HTTPException(status_code=422, detail="O ajuste não pode deixar o estoque negativo.")
        item.quantity_in_stock = new_quantity
        item.unit_cost = unit_cost
        item.min_threshold = min_threshold
        item.category = category.strip()
    else:
        if qty_change < 0:
            raise HTTPException(status_code=422, detail="Não é possível criar item com estoque inicial negativo.")
        db.add(InventoryModel(tenant_id=current_user.tenant_id, part_name=part_name, category=category.strip(), quantity_in_stock=qty_change, min_threshold=min_threshold, unit_cost=unit_cost))
    
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="ATUALIZAR_INVENTARIO", details=f"Item {part_name}; ajuste de quantidade {qty_change}."))
    db.commit()
    return {"status": "success", "message": "Inventário atualizado com sucesso!"}

@app.websocket("/ws/panels")
async def websocket_endpoint(websocket: WebSocket):
    # WebSocket não aceita OAuth2PasswordBearer automaticamente; token via query string.
    token = websocket.query_params.get("token", "")
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        tenant_id = payload.get("tenant_id")
        if not payload.get("sub") or not tenant_id:
            raise jwt.InvalidTokenError("claims ausentes")
    except jwt.PyJWTError:
        await websocket.close(code=1008)
        return
    await manager.connect(websocket, tenant_id)
    try:
        while True:
            await websocket.receive_text()
    except WebSocketDisconnect:
        manager.disconnect(websocket)

@app.get("/api/export/executive-report")
def export_executive_report(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Exporta relatório executivo Excel formatado, limitado ao tenant autenticado."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    output = io.BytesIO()
    def safe_text(value):
        text = "" if value is None else str(value)
        # Evita que conteúdo digitado por usuários seja interpretado como fórmula ao abrir no Excel.
        if text.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + text
        return text

    panels = db.query(PanelModel).filter(PanelModel.tenant_id == current_user.tenant_id).order_by(PanelModel.id).all()
    alarms = db.query(AlarmModel).filter(AlarmModel.tenant_id == current_user.tenant_id).order_by(AlarmModel.first_triggered_at.desc()).all()
    work_orders = db.query(WorkOrderModel).filter(WorkOrderModel.tenant_id == current_user.tenant_id).order_by(WorkOrderModel.id).all()
    assets = db.query(AssetModel).filter(AssetModel.tenant_id == current_user.tenant_id).order_by(AssetModel.id).all()
    tasks = db.query(MaintenanceTaskModel).filter(MaintenanceTaskModel.tenant_id == current_user.tenant_id).order_by(MaintenanceTaskModel.due_at).all()
    inventory = db.query(InventoryModel).filter(InventoryModel.tenant_id == current_user.tenant_id).order_by(InventoryModel.id).all()
    work_order_ids = [item.id for item in work_orders]
    checklist_items = db.query(WorkOrderChecklistItemModel).filter(
        WorkOrderChecklistItemModel.tenant_id == current_user.tenant_id,
        WorkOrderChecklistItemModel.work_order_id.in_(work_order_ids),
    ).order_by(WorkOrderChecklistItemModel.work_order_id, WorkOrderChecklistItemModel.id).all() if work_order_ids else []
    checklist_by_order = {}
    for item in checklist_items:
        checklist_by_order.setdefault(item.work_order_id, []).append(item)
    work_order_by_alarm = {item.alarm_id: item.id for item in work_orders if item.alarm_id is not None}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    dark_green = "064E3B"
    accent = "10B981"
    pale_green = "D1FAE5"
    light_gray = "F3F4F6"
    white = "FFFFFF"
    thin_green = Side(style="thin", color="A7F3D0")

    def setup_sheet(sheet, title, headers, rows):
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = sheet.cell(1, 1, title)
        title_cell.font = Font(name="Calibri", size=16, bold=True, color=white)
        title_cell.fill = PatternFill("solid", fgColor=dark_green)
        title_cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 30
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        sheet.cell(2, 1, f"Organização: {safe_text(current_user.tenant_id)}  |  Gerado em UTC: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')}")
        sheet.cell(2, 1).font = Font(italic=True, color="4B5563", size=10)
        sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[2].height = 24
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
        sheet.cell(3, 1, "Retrato dos dados cadastrados no momento da exportação; dados demonstrativos não substituem avaliação operacional.")
        sheet.cell(3, 1).font = Font(italic=True, color="6B7280", size=9)
        sheet.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[3].height = 28
        header_row = 5
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(header_row, col, header)
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=accent)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=thin_green)
        sheet.row_dimensions[header_row].height = 32
        for row_idx, row in enumerate(rows, header_row + 1):
            for col_idx, value in enumerate(row, 1):
                if isinstance(value, str):
                    value = safe_text(value)
                cell = sheet.cell(row_idx, col_idx, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_idx % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=light_gray)
            sheet.row_dimensions[row_idx].height = 30 if len(headers) > 5 else 21
        last_row = max(header_row, sheet.max_row)
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
        for column in range(1, len(headers) + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(header_row, min(last_row, header_row + 100) + 1)]
            width = min(48, max(12, max((len(value) for value in values), default=12) + 2))
            sheet.column_dimensions[get_column_letter(column)].width = width
        if sheet.title == "Resumo":
            # O título e as notas ocupam células mescladas; considere esses textos na largura.
            sheet.column_dimensions["A"].width = 56
            sheet.column_dimensions["B"].width = 34
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.outlinePr.summaryBelow = True

    summary_rows = [
        ["Painéis", len(panels)],
        ["Painéis críticos", sum(1 for p in panels if p.status == "CRITICO")],
        ["Painéis normais", sum(1 for p in panels if p.status == "NORMAL")],
        ["Alarmes", len(alarms)],
        ["Alarmes ativos", sum(1 for a in alarms if a.status == "ACTIVE")],
        ["Alarmes reconhecidos", sum(1 for a in alarms if a.status == "ACKNOWLEDGED")],
        ["Alarmes encerrados", sum(1 for a in alarms if a.status == "CLEARED")],
        ["Ordens de serviço abertas", sum(1 for w in work_orders if w.status not in ("CONCLUIDA", "CANCELADA"))],
        ["Ordens originadas de alarmes", sum(1 for w in work_orders if w.alarm_id is not None)],
        ["Ativos cadastrados", len(assets)],
        ["Tarefas de manutenção", len(tasks)],
        ["Itens de estoque abaixo do mínimo", sum(1 for i in inventory if i.quantity_in_stock < i.min_threshold)],
    ]
    summary.sheet_view.showGridLines = False
    summary.sheet_view.zoomScale = 90
    summary.sheet_view.selection[0].activeCell = "A1"
    summary.sheet_view.selection[0].sqref = "A1"
    for column in range(1, 13):
        summary.column_dimensions[get_column_letter(column)].width = 18

    def merge_styled(row, start_col, end_col, value, fill_color, font, alignment, border):
        summary.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        for column in range(start_col, end_col + 1):
            cell = summary.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.border = border
        cell = summary.cell(row, start_col)
        cell.value = value
        cell.font = font
        cell.alignment = alignment

    report_border = Border(
        left=thin_green, right=thin_green, top=thin_green, bottom=thin_green,
    )
    merge_styled(
        1, 1, 12, "Relatório Executivo Operacional — HITECH VoltGuard", dark_green,
        Font(name="Calibri", size=20, bold=True, color=white),
        Alignment(vertical="center", horizontal="left", indent=1), report_border,
    )
    summary.row_dimensions[1].height = 42
    merge_styled(
        2, 1, 12,
        f"Organização: {safe_text(current_user.tenant_id)}   |   Gerado em UTC: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')}",
        "ECFDF5", Font(name="Calibri", size=10, bold=True, color=dark_green),
        Alignment(vertical="center", horizontal="left", indent=1), report_border,
    )
    summary.row_dimensions[2].height = 25
    merge_styled(
        3, 1, 12,
        "Retrato dos dados cadastrados no momento da exportação. Dados demonstrativos não substituem avaliação operacional.",
        "F0FDF4", Font(name="Calibri", size=9, italic=True, color="4B5563"),
        Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1), report_border,
    )
    summary.row_dimensions[3].height = 28
    summary.row_dimensions[4].height = 12

    card_alert_colors = {
        "Painéis críticos": "B91C1C",
        "Alarmes ativos": "B91C1C",
        "Alarmes reconhecidos": "B45309",
        "Itens de estoque abaixo do mínimo": "B45309",
    }
    for index, (label, value) in enumerate(summary_rows):
        card_row = 5 + (index // 3) * 3
        start_col = 1 + (index % 3) * 4
        end_col = start_col + 3
        alert_color = card_alert_colors.get(label)
        value_color = alert_color if alert_color and value > 0 else dark_green
        merge_styled(
            card_row, start_col, end_col, label.upper(), value_color,
            Font(name="Calibri", size=12, bold=True, color=white),
            Alignment(vertical="center", horizontal="left", indent=1), report_border,
        )
        merge_styled(
            card_row + 1, start_col, end_col, value, "F0FDF4",
            Font(name="Calibri", size=28, bold=True, color=value_color),
            Alignment(vertical="center", horizontal="left", indent=1), report_border,
        )
        summary.row_dimensions[card_row].height = 27
        summary.row_dimensions[card_row + 1].height = 43
        if index % 3 == 2 and index < len(summary_rows) - 1:
            summary.row_dimensions[card_row + 2].height = 10

    footer_row = 17
    merge_styled(
        footer_row, 1, 12, "Consulte as abas seguintes para os registros detalhados e a rastreabilidade operacional.",
        "ECFDF5", Font(name="Calibri", size=9, italic=True, color="4B5563"),
        Alignment(vertical="center", horizontal="left", indent=1), report_border,
    )
    summary.row_dimensions[footer_row].height = 24
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.page_setup.orientation = "landscape"
    summary.page_margins.left = 0.25
    summary.page_margins.right = 0.25
    summary.page_margins.top = 0.4
    summary.page_margins.bottom = 0.4
    summary.print_area = "A1:L17"
    summary.print_options.horizontalCentered = True

    detail_sheets = [
        ("Painéis", ["ID", "Nome", "Local", "Tensão", "Corrente (A)", "Temperatura (°C)", "Estado", "Componente", "Recomendação"], [[p.id, p.name, p.location, p.voltage, p.current, p.temperature, p.status, p.component, p.recommendation] for p in panels]),
        ("Alarmes", ["ID", "Data inicial", "Última ocorrência", "Reconhecido em", "Reconhecido por", "Encerrado em", "Painel", "Tipo", "Severidade", "Estado", "Mensagem", "OS vinculada"], [[a.id, a.first_triggered_at, a.last_triggered_at, a.acknowledged_at, a.acknowledged_by, a.cleared_at, a.panel_id, a.alarm_type, a.severity, a.status, a.message, work_order_by_alarm.get(a.id)] for a in alarms]),
        ("Ordens de Serviço", ["ID", "Alarme de origem", "Painel", "Título", "Prioridade", "Estado", "Técnico", "Criada em", "Prazo", "Início", "Conclusão", "Checklist concluído", "Itens no checklist", "Resumo de execução"], [[w.id, w.alarm_id, w.panel_id, w.title, w.priority, w.status, w.assigned_technician, w.created_at, w.due_at, w.started_at, w.completed_at, sum(1 for item in checklist_by_order.get(w.id, []) if item.is_completed), len(checklist_by_order.get(w.id, [])), w.resolution] for w in work_orders]),
        ("Ativos", ["ID", "Tag", "Nome", "Tipo", "Local", "Fabricante", "Modelo", "Série", "Painel", "Estado"], [[a.id, a.asset_tag, a.name, a.asset_type, a.location, a.manufacturer, a.model, a.serial_number, a.panel_id, a.status] for a in assets]),
        ("Manutenção", ["ID", "Ativo ID", "Tarefa", "Responsável", "Prazo", "Estado", "Concluída em"], [[t.id, t.asset_id, t.title, t.assigned_to, t.due_at, t.status, t.completed_at] for t in tasks]),
        ("Estoque", ["ID", "Item", "Categoria", "Quantidade", "Estoque mínimo", "Custo unitário (R$)", "Abaixo do mínimo"], [[i.id, i.part_name, i.category, i.quantity_in_stock, i.min_threshold, i.unit_cost, "SIM" if i.quantity_in_stock < i.min_threshold else "NÃO"] for i in inventory]),
    ]
    for sheet_name, headers, rows in detail_sheets:
        setup_sheet(workbook.create_sheet(sheet_name), sheet_name, headers, rows)

    # Start readers on the executive summary and give every sheet a consistent print header.
    workbook.active = 0
    for sheet in workbook.worksheets:
        sheet.print_title_rows = "1:5"
        sheet.sheet_view.zoomScale = 140 if sheet.title == "Resumo" else 90
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(iter([output.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=relatorio_executivo_voltguard.xlsx"})

class WorkOrderCreate(BaseModel):
    panel_id: str
    title: str
    priority: str = "ALTA"
    description: Optional[str] = None
    assigned_technician: Optional[str] = None
    due_at: Optional[datetime] = None
    checklist: Optional[List[str]] = None

class WorkOrderUpdate(BaseModel):
    title: Optional[str] = None
    status: Optional[str] = None
    priority: Optional[str] = None
    assigned_technician: Optional[str] = None
    description: Optional[str] = None
    due_at: Optional[datetime] = None
    resolution: Optional[str] = None

class WorkOrderChecklistItemUpdate(BaseModel):
    is_completed: bool

class WorkOrderChecklistTemplateCreate(BaseModel):
    name: str
    items: List[str]

@app.get("/api/work-order-checklist-templates")
def list_work_order_checklist_templates(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    templates = db.query(WorkOrderChecklistTemplateModel).filter(
        WorkOrderChecklistTemplateModel.tenant_id == current_user.tenant_id
    ).order_by(WorkOrderChecklistTemplateModel.name.asc()).all()
    return [{"id": item.id, "name": item.name, "items": json.loads(item.items_json),
             "created_by": item.created_by, "created_at": utc_iso(item.created_at)} for item in templates]

@app.post("/api/work-order-checklist-templates", status_code=201)
def create_work_order_checklist_template(data: WorkOrderChecklistTemplateCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}:
        raise HTTPException(status_code=403, detail="Somente CEO ou Gerente pode criar modelos de checklist.")
    name = data.name.strip()
    items = [item.strip() for item in data.items if item.strip()]
    if not name or len(name) > 120:
        raise HTTPException(status_code=422, detail="Informe um nome de modelo com até 120 caracteres.")
    if len(data.items) > 20:
        raise HTTPException(status_code=422, detail="O modelo aceita no máximo 20 itens.")
    if not items or len(items) > 20 or any(len(item) > 240 for item in items):
        raise HTTPException(status_code=422, detail="O modelo deve ter de 1 a 20 itens, com até 240 caracteres cada.")
    existing = db.query(WorkOrderChecklistTemplateModel).filter(
        WorkOrderChecklistTemplateModel.tenant_id == current_user.tenant_id,
        WorkOrderChecklistTemplateModel.name == name,
    ).first()
    if existing:
        raise HTTPException(status_code=409, detail="Já existe um modelo com esse nome nesta organização.")
    template = WorkOrderChecklistTemplateModel(
        tenant_id=current_user.tenant_id, name=name, items_json=json.dumps(items, ensure_ascii=False),
        created_by=current_user.username,
    )
    db.add(template)
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username,
                         role=current_user.role, action="CRIAR_MODELO_CHECKLIST_OS",
                         details=f"Modelo de checklist criado: {name} ({len(items)} itens)."))
    db.commit()
    db.refresh(template)
    return {"id": template.id, "name": template.name, "items": items,
            "created_by": template.created_by, "created_at": utc_iso(template.created_at)}

@app.delete("/api/work-order-checklist-templates/{template_id}")
def delete_work_order_checklist_template(template_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}:
        raise HTTPException(status_code=403, detail="Somente CEO ou Gerente pode remover modelos de checklist.")
    template = db.query(WorkOrderChecklistTemplateModel).filter(
        WorkOrderChecklistTemplateModel.id == template_id,
        WorkOrderChecklistTemplateModel.tenant_id == current_user.tenant_id,
    ).first()
    if not template:
        raise HTTPException(status_code=404, detail="Modelo de checklist não encontrado nesta organização.")
    name = template.name
    db.delete(template)
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username,
                         role=current_user.role, action="REMOVER_MODELO_CHECKLIST_OS",
                         details=f"Modelo de checklist removido: {name}."))
    db.commit()
    return {"status": "success", "id": template_id}

@app.post("/api/work-orders", status_code=201)
def create_work_order(data: WorkOrderCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in ["CEO", "Gerente", "Técnico de Campo"]:
        raise HTTPException(status_code=403, detail="Sem permissão para criar ordens de serviço.")
    panel_id, title = data.panel_id.strip(), data.title.strip()
    if not panel_id or len(panel_id) > 100 or not title or len(title) > 200:
        raise HTTPException(status_code=422, detail="ID do painel ou título inválido.")
    if data.priority not in {"BAIXA", "MEDIA", "MÉDIA", "ALTA", "CRITICA", "CRÍTICA"}:
        raise HTTPException(status_code=422, detail="Prioridade inválida.")
    if data.description and len(data.description) > 4000:
        raise HTTPException(status_code=422, detail="Descrição demasiado longa.")
    checklist_labels = [label.strip() for label in (data.checklist or []) if label.strip()]
    if len(checklist_labels) > 20 or any(len(label) > 240 for label in checklist_labels):
        raise HTTPException(status_code=422, detail="O checklist aceita até 20 itens com no máximo 240 caracteres cada.")
    panel = db.query(PanelModel).filter(PanelModel.id == panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel:
        raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    order = WorkOrderModel(tenant_id=current_user.tenant_id, panel_id=panel_id, title=title, priority=data.priority, status="ABERTA", assigned_technician=(data.assigned_technician or "").strip() or None, description=data.description, due_at=utc_naive(data.due_at))
    db.add(order)
    db.flush()
    db.add_all([
        WorkOrderChecklistItemModel(
            tenant_id=current_user.tenant_id, work_order_id=order.id, label=label
        ) for label in checklist_labels
    ])
    db.add(WorkOrderHistoryModel(
        tenant_id=current_user.tenant_id, work_order_id=order.id,
        user_name=current_user.username, role=current_user.role, action="CRIADA",
        details=f"Ordem criada para o painel {panel_id}. Título: {title}. Prioridade: {data.priority}. Técnico: {order.assigned_technician or 'não atribuído'}. Estado inicial: ABERTA.",
    ))
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CRIAR_ORDEM_SERVICO", details=f"Ordem para painel {panel_id}: {title}"))
    db.commit(); db.refresh(order)
    return {"id": order.id, "panel_id": order.panel_id, "title": order.title, "priority": order.priority, "status": order.status, "assigned_technician": order.assigned_technician, "created_at": utc_iso(order.created_at), "due_at": utc_iso(order.due_at), "started_at": utc_iso(order.started_at), "completed_at": utc_iso(order.completed_at), "description": order.description, "resolution": order.resolution}

@app.patch("/api/work-orders/{order_id}")
def update_work_order(order_id: int, data: WorkOrderUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in ["CEO", "Gerente", "Técnico de Campo"]:
        raise HTTPException(status_code=403, detail="Sem permissão para atualizar ordens de serviço.")
    order = db.query(WorkOrderModel).filter(WorkOrderModel.id == order_id, WorkOrderModel.tenant_id == current_user.tenant_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Ordem de serviço não encontrada.")
    provided_fields = getattr(data, "model_fields_set", getattr(data, "__fields_set__", set()))
    previous_status = order.status
    previous_values = {
        "title": order.title,
        "priority": order.priority,
        "assigned_technician": order.assigned_technician,
        "due_at": utc_iso(order.due_at),
        "description": order.description,
        "resolution": order.resolution,
    }
    if data.resolution is not None and len(data.resolution) > 4000:
        raise HTTPException(status_code=422, detail="Resumo de execução demasiado longo.")
    if data.title is not None:
        title = data.title.strip()
        if not title or len(title) > 200:
            raise HTTPException(status_code=422, detail="Informe um título com até 200 caracteres.")
        order.title = title
    allowed_status = {"ABERTA", "EM_ANDAMENTO", "AGUARDANDO_PECA", "CONCLUIDA", "CANCELADA"}
    if data.status is not None:
        if data.status not in allowed_status:
            raise HTTPException(status_code=422, detail="Estado inválido.")
        allowed_transitions = {
            "ABERTA": {"ABERTA", "EM_ANDAMENTO", "AGUARDANDO_PECA", "CONCLUIDA", "CANCELADA"},
            "EM_ANDAMENTO": {"EM_ANDAMENTO", "ABERTA", "AGUARDANDO_PECA", "CONCLUIDA", "CANCELADA"},
            "AGUARDANDO_PECA": {"AGUARDANDO_PECA", "EM_ANDAMENTO", "CANCELADA", "CONCLUIDA"},
            "CONCLUIDA": {"CONCLUIDA"},
            "CANCELADA": {"CANCELADA"},
        }
        if data.status not in allowed_transitions.get(order.status, set()):
            raise HTTPException(status_code=409, detail="Transição inválida para o estado atual da ordem.")
        if data.status == "CONCLUIDA" and previous_status != "CONCLUIDA":
            completion_note = (data.resolution or order.resolution or "").strip()
            if len(completion_note) < 10:
                raise HTTPException(status_code=422, detail="Descreva a solução executada (mínimo de 10 caracteres) antes de concluir a ordem.")
            pending_checklist = db.query(WorkOrderChecklistItemModel).filter(
                WorkOrderChecklistItemModel.tenant_id == current_user.tenant_id,
                WorkOrderChecklistItemModel.work_order_id == order.id,
                WorkOrderChecklistItemModel.is_completed.is_(False),
            ).count()
            if pending_checklist:
                raise HTTPException(status_code=422, detail=f"Conclua os {pending_checklist} item(ns) pendente(s) do checklist antes de encerrar a ordem.")
            order.resolution = completion_note
            order.completed_at = datetime.utcnow()
        if data.status == "EM_ANDAMENTO" and order.started_at is None:
            order.started_at = datetime.utcnow()
        order.status = data.status
    if data.priority is not None:
        if data.priority not in {"BAIXA", "MEDIA", "MÉDIA", "ALTA", "CRITICA", "CRÍTICA"}:
            raise HTTPException(status_code=422, detail="Prioridade inválida.")
        order.priority = data.priority
    if "assigned_technician" in provided_fields:
        technician = (data.assigned_technician or "").strip()
        if len(technician) > 120:
            raise HTTPException(status_code=422, detail="Nome do técnico demasiado longo.")
        order.assigned_technician = technician or None
    if data.description is not None:
        if len(data.description) > 4000:
            raise HTTPException(status_code=422, detail="Descrição demasiado longa.")
        order.description = data.description
    if "due_at" in provided_fields:
        order.due_at = utc_naive(data.due_at)
    if data.resolution is not None:
        order.resolution = data.resolution.strip() or None
    changes = []
    if previous_values["title"] != order.title:
        changes.append(f"Título: {previous_values['title']} → {order.title}")
    if previous_status != order.status:
        changes.append(f"Estado: {previous_status} → {order.status}")
    if previous_values["priority"] != order.priority:
        changes.append(f"Prioridade: {previous_values['priority']} → {order.priority}")
    if previous_values["assigned_technician"] != order.assigned_technician:
        changes.append(f"Técnico: {previous_values['assigned_technician'] or '—'} → {order.assigned_technician or '—'}")
    if previous_values["due_at"] != utc_iso(order.due_at):
        changes.append(f"Prazo: {previous_values['due_at'] or '—'} → {utc_iso(order.due_at) or '—'}")
    if previous_values["description"] != order.description:
        changes.append("Descrição atualizada")
    resolution_changed = previous_values["resolution"] != order.resolution
    if resolution_changed:
        changes.append("Resumo de execução atualizado")
    if changes:
        event_details = "; ".join(changes)
        db.add(WorkOrderHistoryModel(
            tenant_id=current_user.tenant_id, work_order_id=order.id,
            user_name=current_user.username, role=current_user.role,
            action="ATUALIZADA", details=event_details,
            resolution=order.resolution if resolution_changed or order.status == "CONCLUIDA" else None,
        ))
        db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="ATUALIZAR_ORDEM_SERVICO", details=f"Ordem {order_id}: {event_details}."))
    db.commit()
    return {"status": "success", "id": order.id, "order_status": order.status,
            "due_at": utc_iso(order.due_at),
            "started_at": utc_iso(order.started_at),
            "completed_at": utc_iso(order.completed_at),
            "resolution": order.resolution}

@app.patch("/api/work-orders/{order_id}/checklist/{item_id}")
def update_work_order_checklist_item(
    order_id: int, item_id: int, data: WorkOrderChecklistItemUpdate,
    db: Session = Depends(get_db), current_user: User = Depends(get_current_user),
):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para atualizar o checklist da ordem.")
    order = db.query(WorkOrderModel).filter(
        WorkOrderModel.id == order_id, WorkOrderModel.tenant_id == current_user.tenant_id
    ).first()
    if not order:
        raise HTTPException(status_code=404, detail="Ordem de serviço não encontrada.")
    if order.status in {"CONCLUIDA", "CANCELADA"}:
        raise HTTPException(status_code=409, detail="O checklist de uma ordem encerrada não pode ser alterado.")
    item = db.query(WorkOrderChecklistItemModel).filter(
        WorkOrderChecklistItemModel.id == item_id,
        WorkOrderChecklistItemModel.work_order_id == order.id,
        WorkOrderChecklistItemModel.tenant_id == current_user.tenant_id,
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item de checklist não encontrado nesta ordem.")
    event_data = None
    if item.is_completed != data.is_completed:
        item.is_completed = data.is_completed
        item.completed_at = datetime.utcnow() if data.is_completed else None
        item.completed_by = current_user.username if data.is_completed else None
        action = "marcado como concluído" if data.is_completed else "reaberto"
        details = f'Checklist: "{item.label}" {action}.'
        event = WorkOrderHistoryModel(
            tenant_id=current_user.tenant_id, work_order_id=order.id,
            user_name=current_user.username, role=current_user.role,
            action="CHECKLIST_ATUALIZADO", details=details,
        )
        db.add(event)
        db.add(AuditLogModel(
            tenant_id=current_user.tenant_id, user_name=current_user.username,
            role=current_user.role, action="ATUALIZAR_CHECKLIST_ORDEM",
            details=f"Ordem {order.id}: {details}",
        ))
        db.commit()
        db.refresh(item)
        db.refresh(event)
        event_data = {
            "timestamp": utc_iso(event.timestamp), "user_name": event.user_name,
            "role": event.role, "action": event.action, "details": event.details,
            "resolution": None,
        }
    return {
        "id": item.id, "is_completed": item.is_completed,
        "completed_at": utc_iso(item.completed_at), "completed_by": item.completed_by,
        "history_event": event_data,
    }

class AssetCreate(BaseModel):
    asset_tag: str
    name: str
    asset_type: str = "PAINEL_ELETRICO"
    location: str = ""
    manufacturer: str = ""
    model: str = ""
    serial_number: str = ""
    panel_id: Optional[str] = None
    notes: str = ""

class AssetUpdate(BaseModel):
    asset_tag: Optional[str] = None
    name: Optional[str] = None
    asset_type: Optional[str] = None
    location: Optional[str] = None
    manufacturer: Optional[str] = None
    model: Optional[str] = None
    serial_number: Optional[str] = None
    panel_id: Optional[str] = None
    status: Optional[str] = None
    notes: Optional[str] = None

class MaintenancePlanCreate(BaseModel):
    asset_id: int
    title: str
    interval_days: int = 30
    next_due_at: Optional[datetime] = None
    checklist: str = "Inspeção visual;Verificar conexões;Registrar observações"

class MaintenanceTaskCreate(BaseModel):
    asset_id: int
    title: str
    due_at: Optional[datetime] = None
    assigned_to: Optional[str] = None
    notes: str = ""

class MaintenanceTaskUpdate(BaseModel):
    status: str
    assigned_to: Optional[str] = None
    notes: Optional[str] = None
    due_at: Optional[datetime] = None

class DeviceCredentialCreate(BaseModel):
    panel_id: str
    name: str

class DeviceTelemetry(BaseModel):
    panel_id: str
    temperature: Optional[float] = None
    current: Optional[float] = None
    voltage: Optional[str] = None
    voltage_v: Optional[float] = None
    active_power_kw: Optional[float] = None
    active_energy_kwh: Optional[float] = None
    frequency_hz: Optional[float] = None
    power_factor: Optional[float] = None
    gateway_qr: Optional[str] = None

@app.get("/api/assets")
def list_assets(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    rows = db.query(AssetModel).filter(AssetModel.tenant_id == current_user.tenant_id).order_by(AssetModel.name).all()
    return [{"id": x.id, "asset_tag": x.asset_tag, "name": x.name, "asset_type": x.asset_type, "location": x.location, "manufacturer": x.manufacturer, "model": x.model, "serial_number": x.serial_number, "panel_id": x.panel_id, "status": x.status, "notes": x.notes, "created_at": x.created_at.isoformat()} for x in rows]

@app.post("/api/assets", status_code=201)
def create_asset(data: AssetCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para cadastrar ativos.")
    tag, name = data.asset_tag.strip(), data.name.strip()
    if not tag or len(tag) > 80 or not name or len(name) > 160:
        raise HTTPException(status_code=422, detail="Identificação ou nome do ativo inválido.")
    if db.query(AssetModel).filter(AssetModel.tenant_id == current_user.tenant_id, AssetModel.asset_tag == tag).first():
        raise HTTPException(status_code=409, detail="Já existe um ativo com essa identificação.")
    if data.panel_id and not db.query(PanelModel).filter(PanelModel.id == data.panel_id, PanelModel.tenant_id == current_user.tenant_id).first():
        raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    for value, limit, label in [(data.location, 200, "Localização"), (data.manufacturer, 120, "Fabricante"), (data.model, 120, "Modelo"), (data.serial_number, 120, "Número de série"), (data.notes, 2000, "Observações")]:
        if len(value or "") > limit:
            raise HTTPException(status_code=422, detail=f"{label} demasiado longo.")
    row = AssetModel(tenant_id=current_user.tenant_id, asset_tag=tag, name=name, asset_type=data.asset_type[:80], location=data.location.strip(), manufacturer=data.manufacturer.strip(), model=data.model.strip(), serial_number=data.serial_number.strip(), panel_id=data.panel_id, notes=data.notes.strip())
    db.add(row); db.flush()
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CRIAR_ATIVO", details=f"Ativo {tag} cadastrado."))
    db.commit(); db.refresh(row)
    return {"id": row.id, "asset_tag": row.asset_tag, "name": row.name, "asset_type": row.asset_type, "location": row.location, "manufacturer": row.manufacturer, "model": row.model, "serial_number": row.serial_number, "panel_id": row.panel_id, "status": row.status, "notes": row.notes, "created_at": row.created_at.isoformat()}

@app.patch("/api/assets/{asset_id}")
def update_asset(asset_id: int, data: AssetUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}:
        raise HTTPException(status_code=403, detail="Sem permissão para editar ativos.")
    row = db.query(AssetModel).filter(AssetModel.id == asset_id, AssetModel.tenant_id == current_user.tenant_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Ativo não encontrado nesta organização.")
    changes = data.dict(exclude_unset=True)
    if not changes:
        raise HTTPException(status_code=422, detail="Informe ao menos um campo para atualizar.")
    if "asset_tag" in changes:
        changes["asset_tag"] = (changes["asset_tag"] or "").strip()
        if not changes["asset_tag"] or len(changes["asset_tag"]) > 80:
            raise HTTPException(status_code=422, detail="Identificação do ativo inválida.")
        duplicate = db.query(AssetModel).filter(AssetModel.tenant_id == current_user.tenant_id, AssetModel.asset_tag == changes["asset_tag"], AssetModel.id != row.id).first()
        if duplicate:
            raise HTTPException(status_code=409, detail="Já existe outro ativo com essa identificação.")
    if "name" in changes:
        changes["name"] = (changes["name"] or "").strip()
        if not changes["name"] or len(changes["name"]) > 160:
            raise HTTPException(status_code=422, detail="Nome do ativo inválido.")
    if "asset_type" in changes:
        changes["asset_type"] = (changes["asset_type"] or "").strip()
        if not changes["asset_type"] or len(changes["asset_type"]) > 80:
            raise HTTPException(status_code=422, detail="Tipo do ativo inválido.")
    limits = {"location": (200, "Localização"), "manufacturer": (120, "Fabricante"), "model": (120, "Modelo"), "serial_number": (120, "Número de série"), "notes": (2000, "Observações")}
    for field, (limit, label) in limits.items():
        if field in changes:
            changes[field] = (changes[field] or "").strip()
            if len(changes[field]) > limit:
                raise HTTPException(status_code=422, detail=f"{label} demasiado longo.")
    if "panel_id" in changes and changes["panel_id"]:
        changes["panel_id"] = changes["panel_id"].strip()
        if not db.query(PanelModel).filter(PanelModel.id == changes["panel_id"], PanelModel.tenant_id == current_user.tenant_id).first():
            raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    if "status" in changes and changes["status"] not in {"ATIVO", "MANUTENCAO", "INATIVO"}:
        raise HTTPException(status_code=422, detail="Estado do ativo inválido.")
    changed_fields = []
    for field, value in changes.items():
        if getattr(row, field) != value:
            setattr(row, field, value)
            changed_fields.append(field)
    if changed_fields:
        db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="ATUALIZAR_ATIVO", details=f"Ativo {row.asset_tag} atualizado; campos: {', '.join(changed_fields)}."))
        db.commit()
        db.refresh(row)
    return {"id": row.id, "asset_tag": row.asset_tag, "name": row.name, "asset_type": row.asset_type, "location": row.location, "manufacturer": row.manufacturer, "model": row.model, "serial_number": row.serial_number, "panel_id": row.panel_id, "status": row.status, "notes": row.notes, "created_at": row.created_at.isoformat()}

@app.get("/api/maintenance/plans")
def list_maintenance_plans(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    rows = db.query(MaintenancePlanModel).filter(MaintenancePlanModel.tenant_id == current_user.tenant_id).order_by(MaintenancePlanModel.next_due_at).all()
    return [{"id": x.id, "asset_id": x.asset_id, "asset_name": db.query(AssetModel.name).filter(AssetModel.id == x.asset_id, AssetModel.tenant_id == current_user.tenant_id).scalar(), "title": x.title, "checklist": x.checklist, "interval_days": x.interval_days, "next_due_at": x.next_due_at.isoformat(), "active": bool(x.active)} for x in rows]

@app.post("/api/maintenance/plans", status_code=201)
def create_maintenance_plan(data: MaintenancePlanCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para criar planos de manutenção.")
    if not data.title.strip() or len(data.title) > 180 or not 1 <= data.interval_days <= 3650 or len(data.checklist) > 4000:
        raise HTTPException(status_code=422, detail="Título, periodicidade ou checklist inválido.")
    asset = db.query(AssetModel).filter(AssetModel.id == data.asset_id, AssetModel.tenant_id == current_user.tenant_id).first()
    if not asset: raise HTTPException(status_code=404, detail="Ativo não encontrado nesta organização.")
    due = data.next_due_at or (datetime.utcnow() + timedelta(days=data.interval_days))
    row = MaintenancePlanModel(tenant_id=current_user.tenant_id, asset_id=asset.id, title=data.title.strip(), interval_days=data.interval_days, next_due_at=due, checklist=data.checklist.strip())
    db.add(row); db.flush()
    task = MaintenanceTaskModel(tenant_id=current_user.tenant_id, plan_id=row.id, asset_id=asset.id, title=row.title, due_at=due, notes="Checklist: " + row.checklist)
    db.add(task)
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CRIAR_PLANO_MANUTENCAO", details=f"Plano {row.title}; ativo {asset.asset_tag}; próxima execução {due.isoformat()}"))
    db.commit(); db.refresh(row)
    return {"id": row.id, "asset_id": row.asset_id, "title": row.title, "interval_days": row.interval_days, "next_due_at": row.next_due_at.isoformat(), "active": True, "first_task_id": task.id}

@app.get("/api/maintenance/tasks")
def list_maintenance_tasks(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    rows = db.query(MaintenanceTaskModel).filter(MaintenanceTaskModel.tenant_id == current_user.tenant_id).order_by(MaintenanceTaskModel.due_at).all()
    return [{"id": x.id, "plan_id": x.plan_id, "asset_id": x.asset_id, "asset_name": db.query(AssetModel.name).filter(AssetModel.id == x.asset_id, AssetModel.tenant_id == current_user.tenant_id).scalar(), "title": x.title, "due_at": x.due_at.isoformat(), "status": x.status, "assigned_to": x.assigned_to, "notes": x.notes, "completed_at": x.completed_at.isoformat() if x.completed_at else None} for x in rows]

@app.post("/api/maintenance/tasks", status_code=201)
def create_maintenance_task(data: MaintenanceTaskCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para criar tarefas.")
    asset = db.query(AssetModel).filter(AssetModel.id == data.asset_id, AssetModel.tenant_id == current_user.tenant_id).first()
    if not asset: raise HTTPException(status_code=404, detail="Ativo não encontrado nesta organização.")
    if not data.title.strip() or len(data.title) > 180 or len(data.notes) > 4000:
        raise HTTPException(status_code=422, detail="Título ou observações inválidos.")
    row = MaintenanceTaskModel(tenant_id=current_user.tenant_id, asset_id=asset.id, title=data.title.strip(), due_at=data.due_at or datetime.utcnow(), assigned_to=(data.assigned_to or "").strip() or None, notes=data.notes.strip())
    db.add(row); db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CRIAR_TAREFA_MANUTENCAO", details=f"Tarefa {row.title}; ativo {asset.asset_tag}")); db.commit(); db.refresh(row)
    return {"id": row.id, "asset_id": row.asset_id, "title": row.title, "due_at": row.due_at.isoformat(), "status": row.status}

@app.patch("/api/maintenance/tasks/{task_id}")
def update_maintenance_task(task_id: int, data: MaintenanceTaskUpdate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente", "Técnico de Campo"}:
        raise HTTPException(status_code=403, detail="Sem permissão para atualizar tarefas de manutenção.")
    row = db.query(MaintenanceTaskModel).filter(MaintenanceTaskModel.id == task_id, MaintenanceTaskModel.tenant_id == current_user.tenant_id).first()
    if not row: raise HTTPException(status_code=404, detail="Tarefa não encontrada.")
    if data.status not in {"PENDENTE", "EM_ANDAMENTO", "CONCLUIDA", "CANCELADA"}: raise HTTPException(status_code=422, detail="Estado inválido.")
    allowed_transitions = {
        "PENDENTE": {"PENDENTE", "EM_ANDAMENTO", "CONCLUIDA", "CANCELADA"},
        "EM_ANDAMENTO": {"EM_ANDAMENTO", "PENDENTE", "CONCLUIDA", "CANCELADA"},
        "CONCLUIDA": {"CONCLUIDA"},
        "CANCELADA": {"CANCELADA"},
    }
    if data.status not in allowed_transitions.get(row.status, set()):
        raise HTTPException(status_code=409, detail="Transição inválida para o estado atual da tarefa.")
    was_completed = row.status == "CONCLUIDA"
    old_status = row.status
    changed_fields = []
    row.status = data.status
    if old_status != row.status: changed_fields.append(f"estado {old_status} → {row.status}")
    if data.assigned_to is not None:
        assigned_to = data.assigned_to[:120].strip() or None
        if assigned_to != row.assigned_to:
            row.assigned_to = assigned_to
            changed_fields.append("responsável")
    if data.notes is not None and data.notes[:4000] != row.notes:
        row.notes = data.notes[:4000]
        changed_fields.append("observações")
    if data.due_at is not None and data.due_at != row.due_at:
        row.due_at = data.due_at
        changed_fields.append("prazo")
    if data.status == "CONCLUIDA":
        if not was_completed or row.completed_at is None:
            row.completed_at = datetime.utcnow()
    else:
        row.completed_at = None
    if changed_fields:
        db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="ATUALIZAR_TAREFA_MANUTENCAO", details=f"Tarefa {task_id}; " + ", ".join(changed_fields)))
    # When a recurring task is completed, schedule the next task from the plan interval.
    if row.status == "CONCLUIDA" and not was_completed and row.plan_id:
        plan = db.query(MaintenancePlanModel).filter(MaintenancePlanModel.id == row.plan_id, MaintenancePlanModel.tenant_id == current_user.tenant_id, MaintenancePlanModel.active == 1).first()
        if plan:
            next_due = max(plan.next_due_at, datetime.utcnow()) + timedelta(days=plan.interval_days)
            plan.next_due_at = next_due
            db.add(MaintenanceTaskModel(tenant_id=current_user.tenant_id, plan_id=plan.id, asset_id=plan.asset_id, title=plan.title, due_at=next_due, notes="Checklist: " + plan.checklist))
    db.commit()
    return {"status": "success", "id": row.id, "task_status": row.status}

@app.get("/api/devices/credentials")
def list_device_credentials(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    rows = db.query(DeviceCredentialModel).filter(DeviceCredentialModel.tenant_id == current_user.tenant_id).order_by(DeviceCredentialModel.created_at.desc()).all()
    return [{"id": x.id, "panel_id": x.panel_id, "name": x.name, "created_at": x.created_at.isoformat(), "last_used_at": x.last_used_at.isoformat() if x.last_used_at else None, "revoked": bool(x.revoked_at)} for x in rows]

@app.post("/api/devices/credentials", status_code=201)
def create_device_credential(data: DeviceCredentialCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}: raise HTTPException(status_code=403, detail="Somente CEO ou Gerente pode criar credenciais.")
    panel = db.query(PanelModel).filter(PanelModel.id == data.panel_id, PanelModel.tenant_id == current_user.tenant_id).first()
    if not panel: raise HTTPException(status_code=404, detail="Painel não encontrado nesta organização.")
    if not data.name.strip() or len(data.name) > 120: raise HTTPException(status_code=422, detail="Nome da credencial inválido.")
    raw = "vgd_" + secrets.token_urlsafe(32)
    row = DeviceCredentialModel(tenant_id=current_user.tenant_id, panel_id=panel.id, name=data.name.strip(), token_hash=hashlib.sha256(raw.encode()).hexdigest())
    db.add(row); db.flush()
    db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="CRIAR_CREDENCIAL_DISPOSITIVO", details=f"Credencial {row.id} criada para painel {panel.id}.")); db.commit(); db.refresh(row)
    return {"id": row.id, "panel_id": row.panel_id, "name": row.name, "device_token": raw, "warning": "Copie agora; o token não será exibido novamente."}

@app.post("/api/devices/credentials/{credential_id}/revoke")
def revoke_device_credential(credential_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if current_user.role not in {"CEO", "Gerente"}: raise HTTPException(status_code=403, detail="Sem permissão para revogar credenciais.")
    row = db.query(DeviceCredentialModel).filter(DeviceCredentialModel.id == credential_id, DeviceCredentialModel.tenant_id == current_user.tenant_id).first()
    if not row: raise HTTPException(status_code=404, detail="Credencial não encontrada.")
    if row.revoked_at:
        return {"status": "already_revoked", "id": credential_id}
    row.revoked_at = datetime.utcnow(); db.add(AuditLogModel(tenant_id=current_user.tenant_id, user_name=current_user.username, role=current_user.role, action="REVOGAR_CREDENCIAL_DISPOSITIVO", details=f"Credencial {credential_id} revogada.")); db.commit()
    return {"status": "revoked", "id": credential_id}

@app.post("/api/device/telemetry", tags=["Device Ingestion"])
def ingest_device_telemetry(data: DeviceTelemetry, authorization: Optional[str] = Header(default=None), db: Session = Depends(get_db)):
    """Ingestão de telemetria com token exclusivo de dispositivo, armazenado apenas como hash."""
    import math
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Informe Authorization: Bearer <device-token>.")
    raw_token = authorization[7:].strip()
    token_hash = hashlib.sha256(raw_token.encode()).hexdigest()
    credential = db.query(DeviceCredentialModel).filter(DeviceCredentialModel.token_hash == token_hash, DeviceCredentialModel.revoked_at.is_(None)).first()
    if not credential:
        raise HTTPException(status_code=401, detail="Credencial de dispositivo inválida ou revogada.")
    if data.panel_id != credential.panel_id:
        raise HTTPException(status_code=403, detail="Esta credencial não está autorizada para o painel informado.")
    measurements = {
        "temperature": (data.temperature, -50, 250, "Temperatura"),
        "current": (data.current, 0, 100000, "Corrente"),
        "voltage_v": (data.voltage_v, 0, 100000, "Tensão"),
        "active_power_kw": (data.active_power_kw, -1000000, 1000000, "Potência ativa"),
        "active_energy_kwh": (data.active_energy_kwh, 0, 1e15, "Energia ativa"),
        "frequency_hz": (data.frequency_hz, 0, 1000, "Frequência"),
        "power_factor": (data.power_factor, -1, 1, "Fator de potência"),
    }
    if not any(value is not None for value, *_ in measurements.values()):
        raise HTTPException(status_code=422, detail="Envie ao menos uma grandeza de telemetria.")
    for value, minimum, maximum, label in measurements.values():
        if value is not None and (not math.isfinite(value) or not minimum <= value <= maximum):
            raise HTTPException(status_code=422, detail=f"{label} fora do intervalo permitido ({minimum} a {maximum}).")
    if data.voltage is not None and len(data.voltage) > 80: raise HTTPException(status_code=422, detail="Tensão textual demasiado longa.")
    if data.gateway_qr is not None and len(data.gateway_qr) > 160: raise HTTPException(status_code=422, detail="Identificador do gateway demasiado longo.")
    panel = db.query(PanelModel).filter(PanelModel.id == credential.panel_id, PanelModel.tenant_id == credential.tenant_id).first()
    if not panel: raise HTTPException(status_code=404, detail="Painel associado não encontrado.")
    rule = get_effective_alarm_rule(db, credential.tenant_id, panel.id)
    temp_limit, current_limit = rule.max_temperature, rule.max_current
    if data.temperature is not None: panel.temperature = data.temperature
    panel.temperature_measured = data.temperature is not None
    if data.current is not None: panel.current = data.current
    if data.voltage: panel.voltage = data.voltage[:80]
    for field in ("voltage_v", "active_power_kw", "active_energy_kwh", "frequency_hz", "power_factor"):
        value = getattr(data, field)
        if value is not None:
            setattr(panel, field, value)
    if data.voltage_v is not None and data.voltage is None:
        panel.voltage = f"{data.voltage_v:g} V"
    now = datetime.utcnow()
    opened = db.query(AlarmModel).filter(AlarmModel.tenant_id == credential.tenant_id, AlarmModel.panel_id == panel.id, AlarmModel.status.in_(["ACTIVE", "ACKNOWLEDGED"])).all()
    opened_types = {item.alarm_type for item in opened}
    active_conditions = []
    if data.temperature is not None and condition_is_active(data.temperature, temp_limit, rule.temperature_hysteresis, "HIGH_TEMPERATURE" in opened_types): active_conditions.append(("HIGH_TEMPERATURE", f"Temperatura {data.temperature} °C acima do limite configurado de {temp_limit} °C."))
    if data.current is not None and condition_is_active(data.current, current_limit, rule.current_hysteresis, "HIGH_CURRENT" in opened_types): active_conditions.append(("HIGH_CURRENT", f"Corrente {data.current} A acima do limite configurado de {current_limit} A."))
    active_types = {item[0] for item in active_conditions}
    for alarm_type, message in active_conditions:
        existing = next((x for x in opened if x.alarm_type == alarm_type), None)
        if existing: existing.last_triggered_at = now; existing.message = message
        else: db.add(AlarmModel(tenant_id=credential.tenant_id, panel_id=panel.id, alarm_type=alarm_type, severity="CRITICAL", message=message, status="ACTIVE", first_triggered_at=now, last_triggered_at=now))
    for alarm in opened:
        measurement_received = (alarm.alarm_type == "HIGH_TEMPERATURE" and data.temperature is not None) or (alarm.alarm_type == "HIGH_CURRENT" and data.current is not None)
        if measurement_received and alarm.alarm_type not in active_types:
            alarm.status = "CLEARED"; alarm.cleared_at = now
    remaining_open = [alarm for alarm in opened if alarm.status in {"ACTIVE", "ACKNOWLEDGED"}]
    panel.status = "CRITICO" if active_conditions or remaining_open else "NORMAL"
    panel.recommendation = "ALERTA: " + "; ".join(message for _, message in active_conditions) if active_conditions else ("Há alarmes abertos; consulte a Central de Alarmes." if remaining_open else ("Leitura recebida; temperatura e corrente não foram informadas nesta amostra." if data.temperature is None and data.current is None else "Sistema a operar dentro dos parâmetros configurados."))
    credential.last_used_at = now
    db.add(PanelHistoryModel(
        panel_id=panel.id, tenant_id=credential.tenant_id,
        timestamp=now.isoformat(timespec="seconds") + "Z", temperature=panel.temperature,
        temperature_measured=data.temperature is not None,
        current=data.current,
        voltage=data.voltage[:80] if data.voltage is not None else (f"{data.voltage_v:g} V" if data.voltage_v is not None else None),
        voltage_v=data.voltage_v,
        active_power_kw=data.active_power_kw, active_energy_kwh=data.active_energy_kwh,
        frequency_hz=data.frequency_hz, power_factor=data.power_factor,
    ))
    db.commit()
    received_measurements = [field for field, (value, *_rest) in measurements.items() if value is not None]
    if data.voltage is not None:
        received_measurements.append("voltage")
    return {"status": "success", "panel_id": panel.id, "received_at": now.isoformat() + "Z", "active_alarms": len(active_conditions), "measurements_received": received_measurements}

@app.get("/")
def serve_index():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    index_path = os.path.join(base_dir, "index.html")
    if not os.path.exists(index_path):
        index_path = os.path.join(base_dir, "index (3).html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    raise HTTPException(status_code=404, detail=f"index.html não encontrado no caminho: {index_path}")
